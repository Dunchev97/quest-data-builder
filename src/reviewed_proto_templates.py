"""Instantiate audited conf workbooks as the base for new potekha objects."""
from __future__ import annotations

import argparse
import json
from pathlib import Path
import re

from openpyxl import load_workbook

try:
    from .converter_tables import validate
    from .export_xlsx import WorkbookSheet, write_workbook
except ImportError:
    from converter_tables import validate
    from export_xlsx import WorkbookSheet, write_workbook

ROOT = Path(__file__).resolve().parents[1]
CATALOG = ROOT / 'data/fun_interactive_object_templates.json'


def reference(template_id, catalog_path=CATALOG):
    config = json.loads(Path(catalog_path).read_text(encoding='utf-8'))
    try:
        entry = config['reviewed_templates'][template_id]
    except KeyError:
        raise ValueError(f'No reviewed workbook for {template_id}') from None
    path = ROOT / entry['source_file']
    errors = validate(path)
    if errors:
        raise ValueError(f'{path}: {errors}')
    return path, entry


def build_rows(template_id, prefix, catalog_path=CATALOG):
    """Keep real Fun14 inputs; rewrite outputs/links and leave every ID blank.

    This produces a technical starting point. Names, art, locations and balance
    are then filled from the new object's design, without rebuilding the sheet.
    """
    if not re.fullmatch(r'[A-Za-z][A-Za-z0-9_]*', prefix):
        raise ValueError('prefix must be a prototype identifier prefix')
    path, entry = reference(template_id, catalog_path)
    source_prefix = entry['source_prefix']
    wb = load_workbook(path, data_only=False)
    try:
        ws = wb['conf']
        rows = []
        r = 1
        while r <= ws.max_row:
            if ws.cell(r, 1).value not in {'sl', 'ml'}:
                rows.append([c.value for c in ws[r]])
                r += 1
                continue
            start = r
            end = start + 2
            while end <= ws.max_row and any(ws.cell(end, c).value not in (None, '')
                                            for c in range(1, ws.max_column + 1)):
                end += 1
            block = [[ws.cell(rr, c).value for c in range(1, ws.max_column + 1)]
                     for rr in range(start, end)]
            types, keys = block[:2]
            replace_col = None
            c = 3
            while c < len(types):
                typ, key = types[c], keys[c]
                if typ == 'replace' and replace_col is None:
                    replace_col = c
                for data in block[2:]:
                    if key == 'id':
                        data[c] = None
                    elif typ == 'object' and re.fullmatch(r'tasks\.\d+', str(key)) and data[c] == 'identifier':
                        data[c + 1] = None
                    elif re.fullmatch(r'tasks\.\d+\.identifier', str(key)):
                        data[c] = None
                    elif key != 'view_classname':
                        # find belongs to the donor. Only its replacement changes.
                        columns = [c + 1] if typ == 'replace' else [c, c + 1] if typ == 'object' else [c]
                        for cc in columns:
                            if isinstance(data[cc], str):
                                data[cc] = data[cc].replace(source_prefix, prefix)
                c += 2 if typ in {'replace', 'object'} else 1
            for data in block[2:]:
                if isinstance(data[2], str): data[2] = data[2].replace(source_prefix, prefix)
            if prefix != source_prefix:
                if replace_col is None:
                    # Pair columns must immediately follow the last parser field.
                    used = max(i for i, v in enumerate(types) if v not in (None, '')) + 1
                    if types[used - 1] == 'object': used += 1
                    replace_col = used
                    for row in block: row.extend([None] * max(0, used + 2 - len(row)))
                    types[used:used + 2] = ['replace', None]
                    keys[used:used + 2] = ['find', 'replace']
                # ml_replace stops at its first empty cell, independently of
                # longer arrays/objects in the same compact block.
                mapping_row = 2
                while mapping_row < len(block) and block[mapping_row][replace_col] not in (None, ''):
                    mapping_row += 1
                if mapping_row == len(block):
                    block.append([None] * len(block[0]))
                assert block[mapping_row][replace_col + 1] in (None, '')
                block[mapping_row][replace_col:replace_col + 2] = [source_prefix, prefix]
            rows.extend(block)
            rows.append([])
            r = end + 1
        return rows
    finally:
        wb.close()


def create(template_id, prefix, output, *, catalog_path=CATALOG):
    output = Path(output)
    if output.exists():
        raise FileExistsError(f'Preserving existing workbook: {output}')
    rows = build_rows(template_id, prefix, catalog_path)
    # This is an object-authoring copy, not a final quest-pack export.
    result = write_workbook(output, [WorkbookSheet('conf', rows, 'reviewed:' + template_id)])
    errors = validate(output)
    if errors:
        raise ValueError(f'{output}: {errors}')
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('template_id')
    parser.add_argument('--prefix', required=True)
    parser.add_argument('--output', required=True, type=Path)
    args = parser.parse_args()
    print(json.dumps(create(args.template_id, args.prefix, args.output), ensure_ascii=False))


if __name__ == '__main__': main()
