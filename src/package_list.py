"""Create the copy-ready package list for shop_invisible from reviewed XLSX."""
from __future__ import annotations

import argparse
import json
from pathlib import Path, PurePosixPath
import re

from openpyxl import load_workbook


SHOP_DIRECTORY = "d:\\domovoy\\trunk\\server\\prototypes\\context_shop\\"
SHOP_FILE = SHOP_DIRECTORY + "shop_invisible.proto.js"


def filename(campaign_id):
    return f"{campaign_id}_shop_invisible_packages.txt"


def collect_packages(campaign_id, files):
    packages = set()
    for path in sorted(map(Path, files)):
        if path.name.startswith("~$"):
            continue
        wb = load_workbook(path, data_only=False)
        try:
            if "conf" not in wb.sheetnames:
                raise ValueError(f"{path}: missing conf sheet")
            ws = wb["conf"]
            for start in range(1, ws.max_row + 1):
                if ws.cell(start, 1).value not in {"sl", "ml"}:
                    continue
                classname_col = next((c for c in range(4, ws.max_column + 1)
                                      if ws.cell(start + 1, c).value == "classname"
                                      and ws.cell(start, c).value == "string"), None)
                row = start + 2
                while row <= ws.max_row and any(ws.cell(row, c).value not in (None, "")
                                               for c in range(1, ws.max_column + 1)):
                    output = ws.cell(row, 3).value
                    if ws.cell(row, 2).value and isinstance(output, str) and output.startswith("/asset_package/"):
                        if not output.endswith(".proto.js"):
                            raise ValueError(f"{path.name}!C{row}: invalid package output")
                        name = ws.cell(row, classname_col).value if classname_col else None
                        name = name or PurePosixPath(output).name.removesuffix(".proto.js")
                        if not isinstance(name, str) or not re.fullmatch(r"[A-Za-z0-9_]+", name) or not name.startswith(campaign_id + "_"):
                            raise ValueError(f"{path.name}:{row}: invalid or foreign package classname: {name!r}")
                        packages.add(name)
                    row += 1
        finally:
            wb.close()
    return sorted(packages, key=lambda name: [int(p) if p.isdigit() else p for p in re.split(r"(\d+)", name)])


def render(packages):
    return SHOP_DIRECTORY + "\n" + SHOP_FILE + "\n\n" + "".join(
        json.dumps(name, ensure_ascii=False) + ",\n" for name in packages)


def build(campaign_id, tables_dir, *, output_txt=None):
    tables_dir = Path(tables_dir)
    files = [p for p in tables_dir.glob("*.xlsx") if not p.name.startswith("~$")]
    if not files:
        raise ValueError(f"No XLSX tables found: {tables_dir}")
    packages = collect_packages(campaign_id, files)
    text = render(packages)
    output = Path(output_txt) if output_txt else tables_dir / filename(campaign_id)
    output.parent.mkdir(parents=True, exist_ok=True)
    if not output.exists() or output.read_text(encoding="utf-8") != text:
        staged = output.with_name("." + output.name + ".tmp")
        try:
            staged.write_text(text, encoding="utf-8", newline="\n")
            if staged.read_text(encoding="utf-8") != text:
                raise ValueError(f"Package list read-back mismatch: {output}")
            staged.replace(output)
        finally:
            if staged.exists():
                staged.unlink()
    lines = output.read_text(encoding="utf-8").splitlines()
    if lines[:3] != [SHOP_DIRECTORY, SHOP_FILE, ""] or [json.loads(line[:-1]) for line in lines[3:]] != packages:
        raise ValueError(f"Package list smoke check failed: {output}")
    return {"txt": str(output), "packages": len(packages)}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("campaign_id")
    parser.add_argument("--tables-dir", type=Path, required=True)
    parser.add_argument("--output-txt", type=Path)
    args = parser.parse_args(argv)
    try:
        result = build(args.campaign_id, args.tables_dir, output_txt=args.output_txt)
    except (ValueError, OSError) as exc:
        print(str(exc))
        return 1
    print(f"{result['txt']}: {result['packages']} packages")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
