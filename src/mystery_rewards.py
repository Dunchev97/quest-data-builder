"""Mandatory quest reward boxes, their recipes and three permanent prize pots."""
from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path
import re
import shutil
from time import time_ns

try:
    from . import export_xlsx
    from .technical_quest_tables import table_prototypes
except ImportError:
    import export_xlsx
    from technical_quest_tables import table_prototypes


def filename(campaign_id):
    return f"MB_RECIPEMB_POTS_{campaign_id}.xlsx"


def first_recipe_ingredients(campaign_id, tables):
    identifier = f"{campaign_id}_R_1_Recipe"
    values = set()
    for path in tables:
        if path.name.startswith("~$"):
            continue
        for output, fields in table_prototypes(path):
            if output.startswith('/recipe/') and (fields.get('identifier') == identifier or output.endswith('/' + identifier + '.proto.js')):
                value = fields.get('ingredients')
                if not isinstance(value, str) or not value.strip() or value.startswith('='):
                    raise ValueError(f"{path}: {identifier}.ingredients must contain the actual ingredient string")
                values.add(value)
    if len(values) != 1:
        raise ValueError(f"Expected one unambiguous {identifier}.ingredients; found {len(values)}")
    return values.pop()


def block(rows, mode, label, donor, output, fields, campaign_id):
    rows.extend([[None, label],
                 [mode, 'string', 'string'] + [t for t, _, _ in fields] + ['replace', None],
                 [None, 'input', 'output'] + [k for _, k, _ in fields] + ['find', 'replace'],
                 [None, donor, output] + [v for _, _, v in fields] + ['Night_2026', campaign_id], []])


def build_rows(campaign_id, ingredients, conditions=None):
    rows = []
    for n in (1, 2, 3):
        box = f'{campaign_id}_R_MB_{n}'
        block(rows, 'ml', f'Мистери-бокс {n}',
              f'/mystery_box/Night_2026/Night_2026_R_MB_{n}.proto.js',
              f'/mystery_box/{campaign_id}/{box}.proto.js', [
                  ('string', 'classname', box), ('string', 'title', f'Тестовый подарок {n}'),
                  ('string', 'description', 'Используй Волшебный котел на чердачке, чтобы создать'),
                  ('int', 'id', None)], campaign_id)
    for n in (1, 2, 3):
        box = f'{campaign_id}_R_MB_{n}'
        recipe = box + '_Recipe'
        block(rows, 'sl', f'Рецепт мистери-бокса {n}',
              '/recipe/Night_2026/Night_2026_R_MB_1_Recipe.proto.js',
              f'/recipe/{campaign_id}/{recipe}.proto.js', [
                  ('string', 'identifier', recipe), ('string', 'reward', f'asset={box}:1'),
                  ('int', 'position', 1), ('int', 'lifespan', 3600),
                  ('string', 'ingredients', ingredients),
                  ('string', 'conditions', conditions or f'accomplished_or_restarted_group={campaign_id}_Story'),
                  ('string', 'rule_spec.type', 'rule'), ('string', 'rule_spec.title', f'Тестовый подарок {n}'),
                  ('string', 'rule_spec.result_asset', box),
                  ('string', 'rule_spec.description', 'Открыв, можно получить что-то одно!'),
                  ('string', 'rule_spec.window', f'RecipeRule_{campaign_id}_R_MB_Window'),
                  ('int', 'id', None)], campaign_id)
    for n in (1, 2, 3):
        pot = f'{campaign_id}_Pot_Multi_Magic_{n}'
        block(rows, 'sl', f'Горшочек {n}',
              '/pot/Night_2026/Night_2026_Pot_Multi_Magic_1.proto.js',
              f'/pot/{campaign_id}/{pot}.proto.js', [
                  ('string', 'classname', pot), ('string', 'title', f'Тестовый волшебный горшочек {n}'),
                  ('string', 'description', 'Волшебный горшочек для цветочных чудес! Выращивай в нём сразу 5 волшебных цветов!'),
                  ('int', 'price', 35), ('int', 'id', None)], campaign_id)
    return rows


def validate_ingredients(path, campaign_id, tables):
    expected = first_recipe_ingredients(campaign_id, tables)
    found = set()
    for output, fields in table_prototypes(path):
        match = re.search(rf'/{re.escape(campaign_id)}_R_MB_([1-9]\d*)_Recipe\.proto\.js$', output)
        if output.startswith('/recipe/') and match:
            n = int(match.group(1))
            if fields.get('ingredients') != expected:
                raise ValueError(f"{path.name}: R_MB_{n}_Recipe.ingredients differs from R_1_Recipe.ingredients")
            found.add(n)
    if not {1, 2, 3}.issubset(found):
        raise ValueError(f'{path.name}: expected all three R_MB_N_Recipe definitions')


def build(campaign_id, tables_dir, *, output_xlsx=None, conditions=None, force=False):
    folder = Path(tables_dir)
    target = Path(output_xlsx) if output_xlsx else folder / filename(campaign_id)
    if target.exists() and not force:
        raise FileExistsError(f'Preserving manual reward table: {target}; regenerate only intentionally with --force')
    ingredients = first_recipe_ingredients(campaign_id, folder.glob('*.xlsx'))
    result = export_xlsx.write_workbook(target, [export_xlsx.WorkbookSheet(
        'conf', build_rows(campaign_id, ingredients, conditions), 'mystery_rewards')], campaign_id=campaign_id)
    validate_ingredients(target, campaign_id, folder.glob('*.xlsx'))
    return result


def ensure(campaign_id, tables_dir):
    folder = Path(tables_dir)
    target = folder / filename(campaign_id)
    if target.exists():
        validate_ingredients(target, campaign_id, folder.glob('*.xlsx'))
    else:
        build(campaign_id, folder)
    return target


def update_ingredients(campaign_id, tables_dir):
    """Change only MB recipe ingredient cells in an existing manual workbook."""
    from openpyxl import load_workbook
    try:
        from .proto_ids import check_replace_access
    except ImportError:
        from proto_ids import check_replace_access
    folder = Path(tables_dir)
    target = folder / filename(campaign_id)
    expected = first_recipe_ingredients(campaign_id, folder.glob('*.xlsx'))
    original = target.read_bytes()
    wb = load_workbook(target, data_only=False)
    ws = wb['conf']
    changes = {}
    try:
        for r in range(1, ws.max_row + 1):
            if ws.cell(r, 1).value not in {'sl', 'ml'}:
                continue
            col = next((c for c in range(4, ws.max_column + 1) if ws.cell(r+1,c).value == 'ingredients'), None)
            if col is None:
                continue
            end = r + 2
            while end <= ws.max_row and any(ws.cell(end,c).value not in (None,'') for c in range(1,ws.max_column+1)):
                output = ws.cell(end,3).value
                if isinstance(output,str) and output.startswith('/recipe/') and re.search(rf'/{re.escape(campaign_id)}_R_MB_[1-9]\d*_Recipe\.proto\.js$',output):
                    if ws.cell(end,col).value != expected:
                        changes[ws.cell(end,col).coordinate] = expected
                        ws.cell(end,col).value = expected
                end += 1
        if changes:
            check_replace_access(target)
            backup = Path(__file__).resolve().parents[1] / 'output' / ('mb_ingredients_' + datetime.now().strftime('%Y%m%d_%H%M%S_%f'))
            backup.mkdir(parents=True)
            (backup / target.name).write_bytes(original)
            stage = backup / ('updated_' + target.name)
            wb.save(stage)
            validate_ingredients(stage, campaign_id, folder.glob('*.xlsx'))
            check = load_workbook(stage,data_only=False)
            before = load_workbook(backup/target.name,data_only=False)
            try:
                for row in before['conf']:
                    for cell in row:
                        if check['conf'][cell.coordinate].value != changes.get(cell.coordinate,cell.value):
                            raise ValueError(f'Unexpected manual cell change: {cell.coordinate}')
            finally:
                check.close()
                before.close()
            if target.read_bytes() != original:
                raise ValueError(f'Workbook changed while updating: {target}')
            local_stage = target.with_name('.' + target.name + f'.{time_ns()}.tmp')
            try:
                shutil.copyfile(stage,local_stage)
                local_stage.replace(target)
            finally:
                if local_stage.exists():
                    local_stage.unlink()
        validate_ingredients(target,campaign_id,folder.glob('*.xlsx'))
        return changes
    finally:
        wb.close()


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('campaign_id')
    parser.add_argument('--tables-dir', type=Path, required=True)
    parser.add_argument('--output-xlsx', type=Path)
    parser.add_argument('--conditions')
    parser.add_argument('--force', action='store_true')
    parser.add_argument('--update-ingredients',action='store_true',help='Update only ingredient cells in an existing manual table')
    args = parser.parse_args(argv)
    if args.update_ingredients:
        try:
            print(update_ingredients(args.campaign_id,args.tables_dir))
        except (ValueError,OSError) as exc:
            print(str(exc))
            return 1
        return 0
    try:
        result = build(args.campaign_id, args.tables_dir, output_xlsx=args.output_xlsx,
                       conditions=args.conditions, force=args.force)
    except (ValueError, OSError) as exc:
        print(str(exc))
        return 1
    print(result['xlsx'])
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
