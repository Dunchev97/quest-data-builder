"""Build the exact FunCollection_4 template, separately from quest craft objects."""
from __future__ import annotations

import argparse
import copy
import json
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

try:
    from .export_xlsx import WorkbookSheet, write_workbook
except ImportError:
    from export_xlsx import WorkbookSheet, write_workbook

ROOT = Path(__file__).resolve().parents[1]
TEMPLATES = ROOT / "data/fun_interactive_object_templates.json"


def build_rows(selection, campaign_id, templates_path=TEMPLATES):
    if selection["template_id"] != "FunCollection_4":
        raise ValueError("This builder handles exact FunCollection_4 only")
    templates = json.loads(templates_path.read_text(encoding="utf-8"))["templates"]
    template = next(t for t in templates if t["id"] == "FunCollection_4") if isinstance(templates,list) else templates["FunCollection_4"]
    context = dict(selection)
    context.update(prefix=campaign_id, collection_number=4, source_prefix="Fun12", source_location_id=398)
    if len(context["collection_rewards"]) != 10:
        raise ValueError("FunCollection_4 requires exactly ten CL rewards")
    for key in template["required_theme_fields"]:
        if not context.get(key):
            raise ValueError(f"Missing FunCollection_4 theme field: {key}")

    def render(value, extra=None):
        if not isinstance(value,str):
            return value
        values = {**context, **(extra or {})}
        def token(match):
            name, _, default = match.group(1).partition("|")
            if name == "price":  # runtime substitution in the game's button text
                return match.group(0)
            if name == "collection_rewards[index].title":
                return values["reward_title"]
            if name == "negative_effect_titles[effect_suffix]":
                return values["effect_title"]
            if name not in values and "|" not in match.group(1):
                raise ValueError(f"Unresolved template field: {name}")
            return str(values.get(name,default))
        return re.sub(r"\{([^{}]+)\}",token,value)

    base = f"{campaign_id}_FunCollection_4"
    neutral = selection["neutral_classname"]
    def map_state(value):
        return re.sub(re.escape(base)+r"(?!_(?:GR|CL)_)", neutral,value) if isinstance(value,str) else value

    rows, previous = [], None
    for source in template["blocks"]:
        if "same_as" in source:
            block = copy.deepcopy(previous)
            block["title"] = source["title"]
            block["rows"] = [{**block["rows"][0],**source["overrides"]}]
        else:
            block = copy.deepcopy(source)
        previous = copy.deepcopy(block)
        state = block["title"].startswith(("Нейтральное", "Хорошее", "Плохое"))
        types, headers = list(block["types"]), list(block["headers"])
        types[0] = "ml" if state else "sl"
        if block["title"].startswith("Негативные"):
            types = (types + [""]*len(headers))[:len(headers)] + ["replace", ""]
            headers += ["find", "replace"]
            block["rows"][0].update(find="Fun12_FunCollection_4", replace=neutral,
                                    find_2="Fun12", replace_2=campaign_id,
                                    find_3="Fun10", replace_3=campaign_id)
        data = []
        for item in block["rows"]:
            repeat = item.get("repeat", "")
            contexts = [{}]
            if repeat == "collection_rewards[1..10]":
                contexts = [{"index":i,"reward_title":reward["title"]} for i,reward in enumerate(context["collection_rewards"],1)]
            elif repeat == "negative_effect_titles":
                contexts = [{"effect_suffix":suffix,"effect_title":title,
                             "negative_effect_description":context.get("negative_effect_descriptions",{}).get(suffix,context.get("negative_effect_description",""))}
                            for suffix,title in context["negative_effect_titles"].items()]
            for extra in contexts:
                values = {key:render(value,extra) for key,value in item.items()}
                row, counts = [], {}
                for key in headers:
                    counts[key] = counts.get(key,0) + 1
                    lookup = f"{key}_{counts[key]}" if counts[key] > 1 else key
                    value = values.get(lookup,"") if key else ""
                    row.append(map_state(value) if key != "input" else value)
                data.append(row)
        if state:
            # Replacements execute on the donor BEFORE field updates. Long resource
            # names must be replaced before the custom INVITE state family.
            c = headers.index("find")
            mappings = [("Fun12_FunCollection_4_GR_1",f"{base}_GR_1"),
                        ("Fun12_FunCollection_4_CL_",f"{base}_CL_"),
                        ("Fun12_FunCollection_4",neutral),("Fun12",campaign_id)]
            data[0][c:c+2] = list(mappings[0])
            for old,new in mappings[1:]:
                row = [""] * len(headers)
                row[c:c+2] = [old,new]
                data.append(row)
        rows.extend([["",render(block["title"])],types,headers,*data,[]])
    return rows


def build_campaign(campaign_id, pack_id, campaigns_dir=ROOT/"campaigns", output=None, *, force=False):
    manifest_path = campaigns_dir/campaign_id/"fun_interactive_objects.selection.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    selection = next(x for x in manifest["selected_objects"] if x["template_id"] == "FunCollection_4")
    target = output or campaigns_dir/campaign_id/pack_id/"review"/"proto"/f"{campaign_id}_funcollection_4.xlsx"
    if target.exists() and not force:
        raise FileExistsError(f"Preserving manually editable FunCollection table: {target}; use --force to replace it")
    result = write_workbook(target,[WorkbookSheet("conf",build_rows(selection,campaign_id),"fun:FunCollection_4")],campaign_id=campaign_id)
    annotate_manual_settings(target, selection)
    return result


def annotate_manual_settings(target, selection):
    wb = load_workbook(target)
    ws = wb.active
    pending = set(selection.get("pending",[]))
    for r in range(1,ws.max_row+1):
        if ws.cell(r,1).value not in {"sl","ml"}:
            continue
        for c in range(4,ws.max_column+1):
            key = ws.cell(r+1,c).value
            cell = ws.cell(r+2,c)
            note = None
            if key == "conditions" and "tech_quest" in pending:
                note = "Ручная настройка: укажи действующий техквест. Имя по образцу донора; его прототип этим файлом не создаётся."
            elif key == "location_tags" and "location_tag" in pending:
                note = "Ручная настройка: тег локации выпадения корма. Сейчас lawn; согласуй со способом получения и описанием ресурса."
            elif key == "replace" and ws.cell(r+2,c-1).value in (398,"398") and "target_location_id" in pending:
                note = "Ручная настройка: ID места появления Тыковки и мышек. Сейчас 398, как в доноре. Меняй replace; исходный find оставь 398."
            if note:
                cell.fill = PatternFill("solid",fgColor="FFF2CC")
                cell.comment = Comment(note,"Codex")
    wb.save(target)
    wb.close()
    check = load_workbook(target)
    assert check.sheetnames == ["conf"]
    check.close()


def ensure_selected_tables(campaign_id, pack_id, campaigns_dir=ROOT/"campaigns"):
    """Stage 6 must account for selected potekha, while preserving manual XLSX edits."""
    manifest = campaigns_dir/campaign_id/"fun_interactive_objects.selection.json"
    if not manifest.exists():
        return []
    from importlib import import_module
    validator = import_module((__package__ + "." if __package__ else "") + "converter_tables")
    files = []
    for selection in json.loads(manifest.read_text(encoding="utf-8"))["selected_objects"]:
        target = campaigns_dir/campaign_id/pack_id/"review"/"proto"/(campaign_id+"_"+selection["template_id"].lower()+".xlsx")
        if not target.exists():
            if selection["template_id"] != "FunCollection_4":
                raise ValueError(f"Selected FunCollection table is missing: {target}")
            build_campaign(campaign_id,pack_id,campaigns_dir,target)
        errors = validator.validate(target)
        if errors:
            raise ValueError(f"{target}: {errors}")
        files.append(str(target))
    return files


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("campaign_id")
    parser.add_argument("--pack",required=True)
    parser.add_argument("--campaigns-dir",type=Path,default=ROOT/"campaigns")
    parser.add_argument("--output",type=Path)
    parser.add_argument("--force",action="store_true",help="Intentionally replace an existing manually editable XLSX.")
    args = parser.parse_args()
    print(json.dumps(build_campaign(args.campaign_id,args.pack,args.campaigns_dir,args.output,force=args.force),ensure_ascii=False,indent=2))


if __name__ == "__main__":
    main()
