"""The on-disk contract of radmirxan/xlsx-to-json.js (sl/ml only)."""
from __future__ import annotations

import argparse
import json
import math
import re
from pathlib import Path

from openpyxl import load_workbook

try:
    from .quest_naming import normalize_generated_reference
except ImportError:
    from quest_naming import normalize_generated_reference

DIRECTIVES = {"sl", "ml"}
TYPES = {"string", "int", "array", "object", "replace", "ignore"}


def normalize_rows(rows, campaign_id=None):
    """Upgrade old exports without changing authored text or task choices."""
    rows = [list(row) for row in rows]
    for row in rows:
        if row and row[0] in {"temp_01", "temp_02"}:
            row[0] = {"temp_01": "sl", "temp_02": "ml"}[row[0]]
    starts = [i for i, row in enumerate(rows) if row and row[0] in DIRECTIVES]
    for start in starts:
        types, keys = rows[start], rows[start + 1]
        end = start + 2
        while end < len(rows) and any(v not in (None, "") for v in rows[end]):
            end += 1
        data_rows = rows[start + 2:end]
        for data in data_rows:
            for col in range(2,len(data)):
                key = keys[col] if col<len(keys) else None
                if key == "id":
                    data[col] = None
                elif key != "find" and campaign_id:
                    data[col] = normalize_generated_reference(data[col],campaign_id)
            for col,typ in enumerate(types):
                if typ == "object" and col+1<len(data) and data[col] == "id":
                    data[col+1] = None
        inputs = [row[1] for row in data_rows if len(row)>1 and row[1]]
        # Compatibility repair for historical reviewed workbooks. The corrected
        # builders use these paths too; keeping this here preserves manual text.
        aliases = {f"behaviour.{key}": f"behaviour.0.window_spec.{key}" for key in
                   ("title_mark_workbencg", "ingredient_description", "description_1", "description_2")}
        keys[:] = [aliases.get(key,key) for key in keys]
        def append_field(typ, key, value):
            width = max(i+1 for i,v in enumerate(keys) if v not in (None,""))
            types[:] = (types + [None]*width)[:width] + [typ]
            keys[:] = (keys + [None]*width)[:width] + [key]
            for data in data_rows:
                data[:] = (data + [None]*width)[:width] + [value(data)]
        if inputs and all("Fun10_Exchanger" in p for p in inputs) and "replace" not in types:
            campaign = next(row[3].split("_Exchanger",1)[0] for row in data_rows if len(row)>3 and row[3])
            append_field("replace", "find", lambda row:"Fun10")
            append_field(None, "replace", lambda row:campaign)
        if inputs and all(re.search(r"Fun10_(ASK|PER)_1\.proto\.js$",p) for p in inputs) and "view_classname" in keys and "tags.0" not in keys:
            append_field("string", "tags.0", lambda row:re.split(r"_(?:ASK|PER)_",row[3])[0] if row[3] else None)
        for col, typ in enumerate(types):
            key = keys[col] if col < len(keys) else None
            if key == "extra.open_now_btn_window":
                types[col] = "string"
            if typ == "int" and key != "extra.open_now_btn_window":
                for data in rows[start + 2:]:
                    if not any(v not in (None, "") for v in data):
                        break
                    if col < len(data) and isinstance(data[col], str) and data[col].strip():
                        try:
                            value = float(data[col])
                        except ValueError:
                            raise ValueError(f"{key}: expected a number, got {data[col]!r}") from None
                        data[col] = int(value) if value.is_integer() else value
        # Historical path in the XLSX examples is missing a directory.
        for data in rows[start + 2:]:
            if not any(v not in (None, "") for v in data):
                break
            if len(data) > 1 and data[1] == "/quest_group/event/Night_2026_Story.proto.js":
                data[1] = "/quest_group/event/Night_2026/Night_2026_Story.proto.js"
    # ml creates ONE prototype. Each subsequent input starts a fresh block;
    # blank B/C rows remain continuations of that same object/array.
    result = []
    types = keys = None
    seen_input = False
    for row in rows:
        if row and row[0] in DIRECTIVES:
            types, keys, seen_input = row, None, False
        elif types is not None and keys is None:
            keys = row
        elif types and (types[0] == "ml" or "replace" in types) and len(row) > 1 and row[1] not in (None, ""):
            if seen_input:
                result.extend([[], list(types), list(keys)])
            seen_input = True
        if not any(v not in (None, "") for v in row):
            types = keys = None
        result.append(row)
    return result


def prepare_sheet(ws):
    """Merge parser-owned pairs; never discard a non-empty neighbouring cell."""
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value not in DIRECTIVES:
            continue
        for c in range(4, ws.max_column + 1):
            typ = ws.cell(r, c).value
            if typ not in {"object", "replace"}:
                continue
            for rr in ([r] if typ == "replace" else [r, r + 1]):
                if ws.cell(rr, c + 1).value not in (None, ""):
                    raise ValueError(f"{ws.title}!{ws.cell(rr,c+1).coordinate}: cannot merge nonempty cell")
                ws.merge_cells(start_row=rr, end_row=rr, start_column=c, end_column=c + 1)


def validate_sheet(ws, *, campaign_id=None, assigned_ids=False):
    errors = []
    blocks = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value not in (None, "")]
    for r in blocks:
        directive = ws.cell(r, 1).value
        if directive not in DIRECTIVES:
            errors.append(f"row {r}: unsupported directive {directive!r}; only sl/ml")
            continue
        if [ws.cell(r,c).value for c in (2,3)] != ["string", "string"] or [ws.cell(r+1,c).value for c in (2,3)] != ["input", "output"]:
            errors.append(f"row {r}: B/C must be string/string, input/output")
        end = r + 2
        while end <= ws.max_row and any(ws.cell(end,c).value not in (None, "") for c in range(1,ws.max_column+1)):
            end += 1
        data_rows = list(range(r + 2, end))
        if campaign_id:
            for rr in data_rows:
                for cc in range(3,ws.max_column+1):
                    if ws.cell(r+1,cc).value == "find":
                        continue
                    value = ws.cell(rr,cc).value
                    if normalize_generated_reference(value,campaign_id) != value:
                        errors.append(f"{ws.cell(rr,cc).coordinate}: generated character reference or output path uses a legacy format")
        inputs = [rr for rr in data_rows if ws.cell(rr,2).value]
        if not inputs or (directive == "ml" and len(inputs) != 1):
            errors.append(f"row {r}: {directive} has {len(inputs)} input rows (ml requires exactly one)")
        if directive == "sl" and len(inputs)>1 and any(ws.cell(r,c).value == "replace" for c in range(4,ws.max_column+1)):
            errors.append(f"row {r}: sl + replace requires a separate block per input (converter resets replace after the first row)")
        for rr in inputs:
            for c in (2,3):
                value = ws.cell(rr,c).value
                if not isinstance(value,str) or not value.startswith("/") or not value.endswith(".proto.js") or ".." in Path(value).parts:
                    errors.append(f"row {rr}: invalid prototype path in column {c}")
        c, seen, gap = 4, set(), False
        while c <= ws.max_column:
            typ, key = ws.cell(r,c).value, ws.cell(r+1,c).value
            if typ in (None, ""):
                gap = True
                c += 1
                continue
            if gap:
                errors.append(f"row {r}: type after a gap in column {c} will be ignored")
            if typ not in TYPES:
                errors.append(f"row {r}: unsupported type {typ!r}")
            if typ == "object" and directive != "ml":
                errors.append(f"row {r}: object requires ml (the checked sl_object handler is broken)")
            if typ not in {"ignore", "replace"}:
                duplicate_id = key == "id" and all(
                    len({ws.cell(rr,cc).value for cc in range(4,ws.max_column+1)
                         if ws.cell(r+1,cc).value == "id"}) == 1
                    and (assigned_ids or ws.cell(rr,c).value in (None, ""))
                    for rr in data_rows)
                if not key or (key in seen and not duplicate_id):
                    errors.append(f"row {r}: missing or duplicate field {key!r}")
                seen.add(key)
            if re.fullmatch(r"tasks\.\d+", str(key)) and (typ != "object" or directive != "ml"):
                errors.append(f"row {r}: tasks.N requires ml + object")
            if typ == "object" and re.fullmatch(r"tasks\.\d+", str(key)):
                for rr in inputs:
                    if ws.cell(rr,2).value != ws.cell(rr,3).value:
                        errors.append(f"row {rr}: task input/output must both reference the generated quest output")
            if typ in {"object", "replace"}:
                for rr in ([r] if typ == "replace" else [r,r+1]):
                    if not any(m.min_row == rr == m.max_row and m.min_col == c and m.max_col == c+1 for m in ws.merged_cells.ranges):
                        errors.append(f"row {rr}: {typ} pair in column {c} must be merged")
                if typ == "replace" and [ws.cell(r+1,cc).value for cc in (c,c+1)] != ["find", "replace"]:
                    errors.append(f"row {r}: replace requires find/replace headers")
            for rr in data_rows:
                value = ws.cell(rr,c).value
                if key == "id" and value not in (None, "") and not assigned_ids:
                    errors.append(f"row {rr}: id must be blank")
                if typ == "object" and value == "id" and ws.cell(rr,c+1).value not in (None, "") and not assigned_ids:
                    errors.append(f"row {rr}: object id must be blank")
                if typ == "int" and value not in (None, "") and (isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value)):
                    errors.append(f"row {rr}: {key} must be an Excel number")
                if typ == "array" and isinstance(value,str) and value.lstrip().startswith("["):
                    errors.append(f"row {rr}: JSON array in an array cell becomes a nested array; use continuation cells")
            c += 2 if typ in {"object", "replace"} else 1
    for row in ws:
        for cell in row:
            if cell.data_type == "f":
                errors.append(f"{cell.coordinate}: formulas must be materialized before conversion")
            if isinstance(cell.value,str) and ("????" in cell.value or "\ufffd" in cell.value):
                errors.append(f"{cell.coordinate}: broken Unicode")
    return errors


def validate(path, *, assigned_ids=False):
    path = Path(path)
    wb = load_workbook(path, data_only=False)
    errors = []
    if wb.sheetnames != ["conf"]:
        errors.append(f"expected exactly one sheet conf, got {wb.sheetnames}")
    campaign_id = path.parents[2].name if path.parent.name == "proto" and path.parent.parent.name == "review" else None
    for ws in wb:
        errors.extend(f"{ws.title}: {error}" for error in validate_sheet(ws,campaign_id=campaign_id,assigned_ids=assigned_ids))
    wb.close()
    if assigned_ids and not errors:
        try:
            from .proto_ids import assigned_id_errors
        except ImportError:
            from proto_ids import assigned_id_errors
        errors.extend(assigned_id_errors([path]))
    return errors


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx",type=Path)
    parser.add_argument("--assigned-ids",action="store_true",help="Require complete, unique user-assigned proto/task IDs instead of authoring blanks")
    args = parser.parse_args()
    errors = validate(args.xlsx,assigned_ids=args.assigned_ids)
    print(json.dumps({"file":str(args.xlsx),"errors":errors},ensure_ascii=False,indent=2))
    return bool(errors)


if __name__ == "__main__":
    raise SystemExit(main())
