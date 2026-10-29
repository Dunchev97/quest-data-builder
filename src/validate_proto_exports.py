"""Validate single-conf XLSX files and run the actual game converter in isolation."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import shutil
import subprocess

from openpyxl import load_workbook

try:
    from .converter_tables import validate
except ImportError:
    from converter_tables import validate

PROJECT = Path(__file__).resolve().parents[1]


def table_paths(path):
    wb = load_workbook(path, data_only=False)
    ws = wb.active
    result = []
    for r in range(1,ws.max_row+1):
        value = ws.cell(r,2).value
        if isinstance(value,str) and value.startswith('/') and value.endswith('.proto.js'):
            result.append((r,value,ws.cell(r,3).value))
    wb.close()
    return result


def parse_cell(value):
    if not isinstance(value,str):
        return value
    try:
        return json.loads(value)
    except ValueError:
        return value.replace('\r\n','\n')


def expected_updates(path):
    """Record intended field writes independently of the JS serializer."""
    wb = load_workbook(path,data_only=False)
    ws, updates = wb.active, []
    for r in range(1,ws.max_row+1):
        mode = ws.cell(r,1).value
        if mode not in {'sl','ml'}:
            continue
        end = r+2
        while end<=ws.max_row and any(ws.cell(end,c).value not in (None,'') for c in range(1,ws.max_column+1)):
            end += 1
        inputs = [rr for rr in range(r+2,end) if ws.cell(rr,2).value]
        for rr in inputs:
            fields, replacements = {}, []
            for c in range(4,ws.max_column+1):
                typ, key, value = ws.cell(r,c).value, ws.cell(r+1,c).value, ws.cell(rr,c).value
                if typ in {'string','int'} and value is not None:
                    fields[key] = str(value).replace('\r\n','\n') if typ=='string' else value
                elif typ == 'array':
                    if mode=='ml':
                        values = [ws.cell(i,c).value for i in range(rr,end)]
                    else:
                        width = next((m.max_col-m.min_col+1 for m in ws.merged_cells.ranges if m.min_row==r==m.max_row and m.min_col==c),1)
                        values = [ws.cell(rr,i).value for i in range(c,c+width)]
                    fields[key] = [parse_cell(v) for v in values if v is not None]
                elif typ == 'object':
                    fields[key] = {str(ws.cell(i,c).value):parse_cell(ws.cell(i,c+1).value)
                                   for i in range(rr,end) if ws.cell(i,c).value is not None}
                elif typ == 'replace':
                    for i in (range(rr,end) if mode=='ml' else [rr]):
                        old,new = ws.cell(i,c).value,ws.cell(i,c+1).value
                        if old not in (None,'') and new not in (None,''):
                            replacements.append((str(old),str(new)))
            updates.append((ws.cell(rr,3).value,fields,replacements))
    wb.close()
    return updates


def read_field(obj,key):
    for part in key.split('.'):
        obj = obj[int(part)] if isinstance(obj,list) else obj[part]
    return obj


def assert_subset(actual,expected,label):
    if isinstance(expected,dict):
        for key,value in expected.items():
            if value is not None:
                if not isinstance(actual,dict) or key not in actual:
                    raise ValueError(f'{label}.{key}: field lost during conversion')
                assert_subset(actual[key],value,f'{label}.{key}')
    elif actual != expected:
        raise ValueError(f'{label}: expected {expected!r}, got {actual!r}')


def replacement_audit_text(value):
    """Graphic aliases are allowed to keep the donor's view_classname."""
    def mechanics(item):
        if isinstance(item, dict):
            return {key: mechanics(child) for key, child in item.items() if key != 'view_classname'}
        if isinstance(item, list):
            return [mechanics(child) for child in item]
        return item
    return json.dumps(mechanics(value), ensure_ascii=False)


def run(files, converter, prototypes, sandbox, node='node', *, assigned_ids=False):
    files = [Path(p).resolve() for p in files]
    for file in files:
        if file.name.endswith('_resources.xlsx'):
            campaign_id = file.name.removesuffix('_resources.xlsx')
            if file.with_name(campaign_id + '_quests.xlsx') in files:
                try:
                    from .gr_table_review import validate as validate_gr_text
                except ImportError:
                    from gr_table_review import validate as validate_gr_text
                validate_gr_text(campaign_id, file.parent)
    for file in files:
        if file.name.startswith('MB_RECIPEMB_POTS_'):
            try:
                from .mystery_rewards import validate_ingredients
            except ImportError:
                from mystery_rewards import validate_ingredients
            validate_ingredients(file, file.stem.removeprefix('MB_RECIPEMB_POTS_'), files)
    assigned_protos, assigned_tasks = {}, {}
    if assigned_ids:
        try:
            from .proto_ids import read_assigned_ids
        except ImportError:
            from proto_ids import read_assigned_ids
        assigned_protos, assigned_tasks = read_assigned_ids(files)
    for file in files:
        if file.name.endswith('_technical_quests.xlsx'):
            try:
                from .technical_quest_tables import validate_cleanup
            except ImportError:
                from technical_quest_tables import validate_cleanup
            validate_cleanup(file, file.name.removesuffix('_technical_quests.xlsx'), files)
    converter, prototypes, sandbox = [Path(p).resolve() for p in (converter,prototypes,sandbox)]
    if sandbox.exists() and any(sandbox.iterdir()):
        raise ValueError('Use a new empty sandbox directory')
    if sandbox == prototypes or prototypes in sandbox.parents or sandbox in prototypes.parents:
        raise ValueError('Sandbox must be separate from the game prototypes tree')
    sandbox.mkdir(parents=True,exist_ok=True)
    jobs, produced = [], set()
    for file in files:
        file = Path(file).resolve()
        errors = validate(file,assigned_ids=assigned_ids)
        if errors:
            raise ValueError(f'{file}: ' + '; '.join(errors))
        paths = table_paths(file)
        for row,input_path,output_path in paths:
            for value in (input_path,output_path):
                if not isinstance(value,str) or '..' in Path(value).parts:
                    raise ValueError(f'Unsafe prototype path at {file}:{row}')
            if input_path not in produced:
                donor = prototypes / input_path.lstrip('/')
                if not donor.is_file():
                    raise ValueError(f'Missing real donor: {donor} ({file.name}:{row})')
                target = sandbox / input_path.lstrip('/')
                target.parent.mkdir(parents=True,exist_ok=True)
                shutil.copyfile(donor,target)
            produced.add(output_path)
        jobs.append({'file':str(file),'updates_expected':len(paths),'outputs':sorted({p[2] for p in paths})})
    config = sandbox.parent / (sandbox.name + '.run.json')
    config.write_text(json.dumps({'root':str(sandbox),'converter':str(converter),'files':[j['file'] for j in jobs]},ensure_ascii=False),encoding='utf-8')
    command = subprocess.run([str(node),str(PROJECT/'tools/validate_proto_exports.cjs'),str(config)],capture_output=True,text=True,encoding='utf-8')
    if command.returncode:
        raise ValueError(command.stderr or command.stdout)
    results = [json.loads(line) for line in command.stdout.splitlines() if line.strip()]
    for job,result in zip(jobs,results,strict=True):
        if len(result['writes']) != job['updates_expected']:
            raise ValueError(f"{job['file']}: converter performed {len(result['writes'])} updates; expected {job['updates_expected']}")
        job['updates_actual'] = len(result['writes'])
        for output in job['outputs']:
            generated = sandbox / output.lstrip('/')
            if not generated.is_file():
                raise ValueError(f'Output missing: {generated}')
            json.loads(generated.read_text(encoding='utf-8'))
    checked_fields = 0
    for job in jobs:
        for output,fields,replacements in expected_updates(Path(job['file'])):
            actual = json.loads((sandbox/output.lstrip('/')).read_text(encoding='utf-8'))
            for key,value in fields.items():
                try:
                    converted = read_field(actual,key)
                except (KeyError,IndexError,TypeError,ValueError):
                    raise ValueError(f'{output}: field {key} lost during conversion') from None
                assert_subset(converted,value,f'{output}:{key}')
                checked_fields += 1
            text = replacement_audit_text(actual)
            for old,new in replacements:
                if old != new and old not in new and old in text:
                    raise ValueError(f'{output}: donor reference {old!r} survived replacement')
    if assigned_ids:
        actual_tasks = {}
        for output, proto_id in assigned_protos.items():
            actual = json.loads((sandbox/output.lstrip('/')).read_text(encoding='utf-8'))
            if actual.get('id') != proto_id:
                raise ValueError(f'{output}: assigned proto id did not survive conversion')
            if output.startswith('/quest/'):
                for index, task in enumerate(actual.get('tasks', [])):
                    key = (output, f'tasks.{index}')
                    if key not in assigned_tasks or task.get('identifier') != assigned_tasks[key]:
                        raise ValueError(f'{output}: unexpected or incorrectly assigned task {index}')
                    actual_tasks[key] = task['identifier']
        if actual_tasks != assigned_tasks:
            raise ValueError('Assigned task identifiers were lost during conversion')
    return {'status':'passed','converter':str(converter),'converter_sha256':hashlib.sha256(converter.read_bytes()).hexdigest(),
            'prototypes':str(prototypes),'sandbox':str(sandbox),'files':jobs,'prototypes_created':len(produced),
            'fields_verified':checked_fields,
            'ids_assigned':assigned_ids,
            'note':'Complete proto/task IDs checked after conversion.' if assigned_ids else 'Donor IDs are inherited when an id cell is blank. Blank does not clear the donor id; assigning production IDs is a separate developer operation.'}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('xlsx',type=Path,nargs='+')
    parser.add_argument('--converter',type=Path,required=True)
    parser.add_argument('--prototypes',type=Path,required=True)
    parser.add_argument('--sandbox',type=Path,required=True)
    parser.add_argument('--node',default='node')
    parser.add_argument('--assigned-ids',action='store_true',help='Validate complete user-assigned proto/task IDs across all files')
    parser.add_argument('--report',type=Path,help='Optional diagnostic JSON outside the final table directory.')
    args = parser.parse_args()
    try:
        result = run(args.xlsx,args.converter,args.prototypes,args.sandbox,args.node,assigned_ids=args.assigned_ids)
    except (ValueError,OSError) as exc:
        print(str(exc))
        return 1
    if args.report:
        args.report.parent.mkdir(parents=True,exist_ok=True)
        args.report.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(f"OK: {len(result['files'])} workbooks; {result['prototypes_created']} prototypes; {result['fields_verified']} fields")
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
