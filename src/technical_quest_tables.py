"""Build the separate technical quest table from the current reviewed XLSX files."""
from __future__ import annotations

import argparse
import json
from pathlib import Path, PurePosixPath

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Border, Font, PatternFill

try:
    from . import converter_tables, export_xlsx
    from .campaigns import DEFAULT_CAMPAIGNS_DIR, pack_dir
except ImportError:
    import converter_tables
    import export_xlsx
    from campaigns import DEFAULT_CAMPAIGNS_DIR, pack_dir


REFERENCE = "Night_2026"
GROUP_DONOR = "/quest_group/event/Night_2026/Night_2026_Tech.proto.js"
QUEST_DONORS = {
    n: f"/quest/event/Night_2026/Night_2026_Tech_{n}.proto.js" for n in (1, 2)
}
ASSET_FAMILIES = {"quest_item", "workbench_item", "mystery_box"}
STUFF_FAMILIES = {"furniture", "debris", "chest", "workbench"}
PERSISTENT_REWARD_FAMILIES = {"pot"}
SERVICE_FAMILIES = {
    "asset_package", "global_reward", "recipe", "post_action", "friend_action",
    "quest_action", "quest", "quest_group", "user_interface",
}


def workbook_name(campaign_id):
    return f"{campaign_id}_technical_quests.xlsx"


def table_prototypes(path):
    """Read explicit prototype names; never collect unrelated ingredient references."""
    wb = load_workbook(path, data_only=False)
    try:
        if wb.sheetnames != ["conf"]:
            raise ValueError(f"{path}: expected exactly one conf sheet")
        ws = wb.active
        for r in range(1, ws.max_row + 1):
            if ws.cell(r, 1).value not in {"sl", "ml"}:
                continue
            keys = {ws.cell(r + 1, c).value: c for c in range(4, ws.max_column + 1)
                    if ws.cell(r, c).value in {"string", "int"}}
            end = r + 2
            while end <= ws.max_row and any(ws.cell(end, c).value not in (None, "")
                                           for c in range(1, ws.max_column + 1)):
                donor, output = ws.cell(end, 2).value, ws.cell(end, 3).value
                if donor:
                    if not isinstance(output, str) or not output.endswith(".proto.js"):
                        raise ValueError(f"{path}:{end}: missing prototype output")
                    fields = {key: ws.cell(end, c).value for key, c in keys.items()}
                    yield output, fields
                end += 1
    finally:
        wb.close()


def collect_cleanup(campaign_id, tables):
    """Include every owned resource/state and deduplicate campaign-wide tables."""
    assets, stuff, crafting = set(), set(), set()
    title = None
    for path in sorted(tables):
        if path.name.startswith("~$") or path.name == workbook_name(campaign_id):
            continue
        for output, fields in table_prototypes(path):
            parts = PurePosixPath(output).parts
            family = parts[1] if parts and parts[0] == "/" and len(parts) > 2 else ""
            name = fields.get("classname") or PurePosixPath(output).name.removesuffix(".proto.js")
            if family not in ASSET_FAMILIES | STUFF_FAMILIES | SERVICE_FAMILIES | PERSISTENT_REWARD_FAMILIES:
                raise ValueError(f"{path.name}: unclassified prototype family {family!r}: {output}")
            if family in ASSET_FAMILIES | STUFF_FAMILIES:
                if not isinstance(name, str) or not name.startswith(campaign_id + "_"):
                    raise ValueError(f"{path.name}: refusing cleanup of a foreign classname: {name!r}")
                (assets if family in ASSET_FAMILIES else stuff).add(name)
            if family == "quest_group" and output.endswith(f"/{campaign_id}_Story.proto.js"):
                title = title or fields.get("title")
            # Standard story recipes use the same cauldron as the verified donor.
            if output.startswith(f"/recipe/{campaign_id}/"):
                crafting.add("DayOfDomovoyCaldron")
            if family == "workbench":
                crafting.add(name)
    overlap = assets & stuff
    if overlap:
        raise ValueError(f"Classnames used as both assets and stuff: {sorted(overlap)}")
    # Cancel before removing resources, so interrupted crafting cannot restore them later.
    actions = [{"do": "cancel_crafting", "workbench": name, "recipe_tags": [campaign_id]}
               for name in sorted(crafting)]
    actions.extend({"do": "remove_asset", "classname": name} for name in sorted(assets))
    actions.extend({"do": "remove_stuff", "classname": name} for name in sorted(stuff))
    return actions, {"assets": len(assets), "stuff": len(stuff), "crafting": len(crafting),
                     "title": title or campaign_id}


def validate_cleanup(path, campaign_id, tables):
    """Catch stale cleanup lists when the user adds or renames a reviewed entity."""
    expected, _ = collect_cleanup(campaign_id, tables)
    wb = load_workbook(path, data_only=False)
    actual = []
    try:
        ws = wb["conf"]
        for r in range(1, ws.max_row + 1):
            if ws.cell(r, 1).value != "ml":
                continue
            for c in range(4, ws.max_column + 1):
                if ws.cell(r, c).value != "array" or ws.cell(r + 1, c).value != "on_spoil":
                    continue
                end = r + 2
                while end <= ws.max_row and ws.cell(end, c).value not in (None, ""):
                    actual.append(json.loads(ws.cell(end, c).value))
                    end += 1
    finally:
        wb.close()
    canonical = lambda action: json.dumps(action, sort_keys=True, ensure_ascii=False)
    present = {canonical(a) for a in actual}
    missing = [a for a in expected if canonical(a) not in present]
    if missing:
        raise ValueError(f"{path.name}: on_spoil missing campaign cleanup: {missing}")
    removed = [a.get("classname") for a in actual if a.get("do") in {"remove_asset", "remove_stuff"}]
    if len(removed) != len(set(removed)):
        raise ValueError(f"{path.name}: duplicate on_spoil classname")
    if any(not isinstance(name, str) or not name.startswith(campaign_id + "_") for name in removed):
        raise ValueError(f"{path.name}: foreign or missing on_spoil classname")


def append_block(rows, label, donor, output, fields, *, replace=None):
    """An ml array cell is ONE JSON element; later rows have blank input/output."""
    columns = [(typ, key, value) for typ, key, value in fields]
    rows.extend([[None, label], ["ml", "string", "string"] + [t for t, _, _ in columns]
                 + (["replace", None] if replace else []),
                 [None, "input", "output"] + [k for _, k, _ in columns]
                 + (["find", "replace"] if replace else [])])
    length = max([1] + [len(v) for t, _, v in columns if t == "array"])
    for i in range(length):
        row = [None, donor if i == 0 else None, output if i == 0 else None]
        for typ, _, value in columns:
            item = value[i] if typ == "array" and i < len(value) else (value if i == 0 and typ != "array" else None)
            row.append(json.dumps(item, ensure_ascii=False, separators=(",", ":"))
                       if isinstance(item, (dict, list)) else item)
        if replace:
            row.extend([REFERENCE, replace] if i == 0 else [None, None])
        rows.append(row)
    rows.append([])


def append_task(rows, output, fields):
    rows.extend([[None, "Задача: identifier заполняется после проверки таблиц"],
                 ["ml", "string", "string", "object", None],
                 [None, "input", "output", "tasks.0", None]])
    for i, (key, value) in enumerate([("identifier", None)] + list(fields)):
        rows.append([None, output if i == 0 else None, output if i == 0 else None, key, value])
    rows.append([])


def build_rows(campaign_id, cleanup, title, spoil_conditions=None):
    group = f"{campaign_id}_Tech"
    quests = {n: f"{campaign_id}_Tech_{n}" for n in (1, 2)}
    outputs = {n: f"/quest/event/{campaign_id}/{quests[n]}.proto.js" for n in (1, 2)}
    rows = []
    append_block(rows, "Техническая квест-группа", GROUP_DONOR,
                 f"/quest_group/event/{campaign_id}/{group}.proto.js", [
                     ("string", "identifier", group), ("string", "title", f"{title} (Технический)"),
                     ("string", "spoil_conditions", spoil_conditions),
                     ("string", "first_quest", quests[1]), ("string", "last_quest", quests[2]),
                     ("int", "id", None)], replace=campaign_id)
    condition = f"active_quest_group={campaign_id}_Story"
    append_block(rows, "Технический квест 1: запуск", QUEST_DONORS[1], outputs[1], [
        ("string", "identifier", quests[1]), ("string", "group_identifier", group),
        ("string", "conditions", condition), ("string", "conditions_my_mail_test", condition),
        ("string", "conditions_my_mail_test_prod", condition),
        ("array", "opens", [quests[2]]), ("int", "id", None)], replace=campaign_id)
    append_task(rows, outputs[1], [("type", "have_location"), ("location_index", 1)])
    append_block(rows, "Технический квест 2: очистка при завершении события", QUEST_DONORS[2], outputs[2], [
        ("string", "identifier", quests[2]), ("string", "group_identifier", group),
        ("array", "on_spoil", cleanup), ("int", "id", None)], replace=campaign_id)
    append_task(rows, outputs[2], [("type", "action"), ("action", "no_such_action"), ("amount", 1)])
    return rows


def build(campaign_id, pack_id, campaigns_dir=DEFAULT_CAMPAIGNS_DIR, *, output_xlsx=None,
          force=False, spoil_conditions=None):
    campaign = Path(campaigns_dir) / campaign_id
    folder = pack_dir(campaign_id, pack_id, campaigns_dir) / "review" / "proto"
    output = Path(output_xlsx) if output_xlsx else folder / workbook_name(campaign_id)
    if output.exists() and not force:
        raise FileExistsError(f"Preserving manually editable technical table: {output}; use --force intentionally")
    tables = sorted(p for p in campaign.glob("pack_*/review/proto/*.xlsx")
                    if not p.name.startswith("~$") and p.name != workbook_name(campaign_id))
    if not tables:
        raise ValueError(f"No reviewed XLSX tables found under {campaign}")
    cleanup, summary = collect_cleanup(campaign_id, tables)
    rows = build_rows(campaign_id, cleanup, summary["title"], spoil_conditions)
    export_xlsx.write_workbook(output, [export_xlsx.WorkbookSheet("conf", rows, "technical_quests")],
                               campaign_id=campaign_id)
    wb = load_workbook(output, data_only=False)
    ws = wb.active
    yellow = PatternFill("solid", fgColor="FFF2CC")
    for row in ws:
        for cell in row:
            if cell.value == "identifier" and cell.column == 4 and ws.cell(cell.row - 2, 4).value == "object":
                target = ws.cell(cell.row, 5)
                target.fill = yellow
                target.comment = Comment("Уникальный identifier задачи пользователь задаст после проверки таблиц.", "Codex")
            if cell.value == "spoil_conditions":
                target = ws.cell(cell.row + 1, cell.column)
                target.fill = yellow
                target.comment = Comment("Задайте срок события: time>YYYY-MM-DD HH:MM. Пустая ячейка наследует устаревшую дату донора; до импорта обязательно заполнить.", "Codex")
    ws.column_dimensions["B"].width = 64
    ws.column_dimensions["C"].width = 64
    ws.column_dimensions["F"].width = 90
    for row in ws:
        if row[0].value == "ml":
            continue
        if any(isinstance(c.value, str) and c.value.startswith('{"do":') for c in row):
            ws.row_dimensions[row[0].row].height = 30
            # Generic one-cell section styling does not apply to array continuations.
            for cell in row:
                cell.fill = PatternFill(fill_type=None)
                cell.font = Font(name="Calibri", size=11, color="1F2933")
                cell.border = Border()
    wb.save(output)
    wb.close()
    errors = converter_tables.validate(output)
    if errors:
        raise ValueError(f"{output}: {errors}")
    with_reopen = load_workbook(output, data_only=False)
    assert with_reopen.sheetnames == ["conf"]
    with_reopen.close()
    return {"xlsx": str(output), **summary}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("campaign_id")
    parser.add_argument("--pack", default="pack_001")
    parser.add_argument("--campaigns-dir", type=Path, default=DEFAULT_CAMPAIGNS_DIR)
    parser.add_argument("--output-xlsx", type=Path)
    parser.add_argument("--spoil-conditions", help="Verified event deadline, e.g. time>2026-11-01 23:59")
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args(argv)
    try:
        result = build(args.campaign_id, args.pack, args.campaigns_dir, output_xlsx=args.output_xlsx,
                       force=args.force, spoil_conditions=args.spoil_conditions)
    except (ValueError, OSError) as exc:
        print(str(exc))
        return 1
    print(f"{result['xlsx']}: assets={result['assets']}, stuff={result['stuff']}, crafting={result['crafting']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
