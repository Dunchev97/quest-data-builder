"""Create a copy-ready DayOfDomovoyCaldron set entry from reviewed XLSX."""
from __future__ import annotations

import argparse
import json
from pathlib import Path, PurePosixPath
import re

from openpyxl import load_workbook


WORKBENCH_FILE = "d:\\domovoy\\trunk\\server\\prototypes\\workbench\\DayOfDomovoyCaldron.proto.js"


def filename(campaign_id):
    return f"{campaign_id}_DayOfDomovoyCaldron.txt"


def natural_key(name):
    return [int(part) if part.isdigit() else part for part in re.split(r"(\d+)", name)]


def collect_recipes(campaign_id, files):
    """Collect definitions from every table, including manually added recipes."""
    recipes = set()
    title = None
    for path in sorted(map(Path, files)):
        if path.name.startswith("~$"):
            continue
        wb = load_workbook(path, data_only=False)
        try:
            # Read legacy manual Conf too; this does not certify converter readiness.
            sheets = [ws for ws in wb if ws.title.casefold() == "conf"]
            if len(sheets) != 1:
                raise ValueError(f"{path}: expected one conf sheet")
            ws = sheets[0]
            for start in range(1, ws.max_row + 1):
                if ws.cell(start, 1).value not in {"sl", "ml"}:
                    continue
                columns = {ws.cell(start + 1, c).value: c
                           for c in range(4, ws.max_column + 1)
                           if ws.cell(start, c).value == "string"}
                row = start + 2
                while row <= ws.max_row and any(ws.cell(row, c).value not in (None, "")
                                               for c in range(1, ws.max_column + 1)):
                    if ws.cell(row, 1).value in {"sl", "ml"}:
                        break
                    output = ws.cell(row, 3).value
                    if ws.cell(row, 2).value and isinstance(output, str):
                        output = output.replace("\\", "/")
                        if output.startswith("/recipe/"):
                            if not output.endswith(".proto.js"):
                                raise ValueError(f"{path.name}!C{row}: invalid recipe output")
                            col = columns.get("classname")
                            name = ws.cell(row, col).value if col else None
                            name = name or PurePosixPath(output).name.removesuffix(".proto.js")
                            if (not isinstance(name, str)
                                    or not re.fullmatch(r"[A-Za-z0-9_]+", name)
                                    or not name.startswith(campaign_id + "_")):
                                raise ValueError(f"{path.name}:{row}: invalid or foreign recipe classname: {name!r}")
                            recipes.add(name)
                        if (output.startswith("/quest_group/")
                                and output.endswith(f"/{campaign_id}_Story.proto.js")
                                and "title" in columns):
                            value = ws.cell(row, columns["title"]).value
                            if isinstance(value, str) and value and not value.startswith("="):
                                title = title or value
                    row += 1
        finally:
            wb.close()
    return sorted(recipes, key=natural_key), title


def make_entry(campaign_id, recipes, title=None):
    first = f"{campaign_id}_R_1_Recipe"
    icon_recipe = first if first in recipes else recipes[0] if recipes else None
    return {
        "id": campaign_id,
        "attention": "Изготавливать можно",
        "title": title or campaign_id,
        "icon": icon_recipe.removesuffix("_Recipe") if icon_recipe else f"{campaign_id}_R_1",
        "conditions": f"active_quest_group={campaign_id}_Tech",
        "recipes": recipes,
    }


def build(campaign_id, tables_dir, *, title=None):
    tables_dir = Path(tables_dir)
    files = [p for p in tables_dir.glob("*.xlsx") if not p.name.startswith("~$")]
    if not files:
        raise ValueError(f"No XLSX tables found: {tables_dir}")
    recipes, detected_title = collect_recipes(campaign_id, files)
    entry = make_entry(campaign_id, recipes, title or detected_title)
    # The trailing comma matches an entry pasted into extra.set, not a whole proto.
    text = WORKBENCH_FILE + "\n\n" + json.dumps(entry, ensure_ascii=False, indent=2) + ",\n"
    output = tables_dir / filename(campaign_id)
    if not output.exists() or output.read_text(encoding="utf-8") != text:
        staged = output.with_name("." + output.name + ".tmp")
        try:
            staged.write_text(text, encoding="utf-8", newline="\n")
            if staged.read_text(encoding="utf-8") != text:
                raise ValueError(f"Workbench addition read-back mismatch: {output}")
            staged.replace(output)
        finally:
            if staged.exists():
                staged.unlink()
    header, payload = output.read_text(encoding="utf-8").split("\n\n", 1)
    if header != WORKBENCH_FILE or json.loads(payload.rstrip().removesuffix(",")) != entry:
        raise ValueError(f"Workbench addition smoke check failed: {output}")
    return {"txt": str(output), "recipes": len(recipes)}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("campaign_id")
    parser.add_argument("--tables-dir", type=Path, required=True)
    parser.add_argument("--title")
    args = parser.parse_args(argv)
    try:
        result = build(args.campaign_id, args.tables_dir, title=args.title)
    except (ValueError, OSError) as exc:
        print(str(exc))
        return 1
    print(f"{result['txt']}: {result['recipes']} recipes")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
