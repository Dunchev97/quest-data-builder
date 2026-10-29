"""Stage 6.1: attach reviewed dialogs to existing conf tables without regenerating them."""
from __future__ import annotations

import copy
from dataclasses import dataclass
import importlib.util
import json
import os
from pathlib import Path
import re
import tempfile
import uuid

from openpyxl import load_workbook
from openpyxl.styles import Alignment

try:
    from . import proto_ids, validate_proto_exports
except ImportError:
    import proto_ids
    import validate_proto_exports

PROJECT = Path(__file__).resolve().parents[1]
REGISTRY = 'dialogue_characters.json'
PLAN = 'dialogue_plan.json'


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def json_bytes(value):
    return (json.dumps(value, ensure_ascii=False, indent=2) + '\n').encode('utf-8')


def sheet_snapshot(ws):
    return ([(c.coordinate, c.value, copy.copy(c._style), copy.copy(c.comment))
             for row in ws for c in row],
            sorted(str(m) for m in ws.merged_cells.ranges), ws.freeze_panes,
            ws.auto_filter.ref, ws.print_area,
            {r: tuple(sorted(dict(d).items())) for r, d in ws.row_dimensions.items()})


def require(condition, message):
    if not condition:
        raise ValueError(message)


def identifier(value):
    require(isinstance(value, str) and re.fullmatch(r'[A-Za-z0-9_]+', value),
            f'Invalid prototype identifier: {value!r}')
    return value


def proto_path(root, relative):
    require(isinstance(relative, str) and relative.startswith('/') and relative.endswith('.proto.js'),
            f'Invalid prototype path: {relative!r}')
    target = (root / relative.lstrip('/')).resolve()
    require(target.is_relative_to(root.resolve()), f'Path escapes prototypes: {relative}')
    return target


def load_registry(path):
    data = read_json(path)
    require(data.get('version') == 1, 'Character registry requires version=1')
    characters = data.get('characters')
    require(isinstance(characters, dict) and characters, 'Character registry is empty')
    classnames = []
    for key, character in characters.items():
        identifier(key)
        for field in ('name', 'classname', 'view_class'):
            require(isinstance(character.get(field), str) and character[field].strip(),
                    f'Missing characters.{key}.{field}')
        identifier(character['classname'])
        identifier(character['view_class'])
        for emotion, view in character.get('emotions', {}).items():
            require(bool(emotion), f'Empty emotion for {key}')
            identifier(view)
        classnames.append(character['classname'])
    require(len(classnames) == len(set(classnames)), 'Duplicate character classname')
    return data


def helper_views(registry):
    result = {}
    for key, char in registry['characters'].items():
        for alias in (key, char['name'], char['classname'], char['view_class']):
            require(alias not in result or result[alias] == char['view_class'],
                    f'Ambiguous helper alias: {alias}')
            result[alias] = char['view_class']
    return result


def map_helper_rows(rows, registry):
    """Resolve UI portraits only; task icons and physical character references stay intact."""
    mapping = helper_views(registry)
    for index, row in enumerate(rows[:-1]):
        if 'helper' not in row:
            continue
        col = row.index('helper')
        value = rows[index + 1][col]
        if value:
            require(value in mapping, f'Unknown quest helper {value!r}; complete {REGISTRY}')
            rows[index + 1][col] = mapping[value]


def scan_quests(ws):
    quests = {}
    for row in range(1, ws.max_row):
        if ws.cell(row, 1).value == 'sl':
            output = ws.cell(row + 2, 3).value
            if isinstance(output, str) and output.startswith('/quest/'):
                quests[Path(output).name.removesuffix('.proto.js')] = {
                    'output': output, 'row': row, 'tasks': []}
        if ws.cell(row, 1).value != 'ml':
            continue
        key = ws.cell(row + 1, 4).value
        if not isinstance(key, str) or not re.fullmatch(r'tasks\.\d+', key):
            continue
        output = ws.cell(row + 2, 3).value
        quest = Path(output).name.removesuffix('.proto.js')
        require(quest in quests, f'Task has no preceding quest block: {output}')
        end, fields = row + 2, {}
        while end <= ws.max_row and ws.cell(end, 4).value is not None:
            fields[str(ws.cell(end, 4).value)] = ws.cell(end, 5).value
            end += 1
        quests[quest]['tasks'].append({'row': row, 'end': end, 'key': key, 'fields': fields})
    require(bool(quests), 'No quests found in conf')
    for quest, spec in quests.items():
        require([t['key'] for t in spec['tasks']] == [f'tasks.{i}' for i in range(len(spec['tasks']))],
                f'Noncontiguous tasks in {quest}')
    return quests


def insert_rows(ws, row, count):
    merges = []
    for merged in list(ws.merged_cells.ranges):
        if merged.max_row < row:
            continue
        require(merged.min_row >= row, f'Insertion crosses merged range {merged}')
        merges.append((merged.min_row, merged.max_row, merged.min_col, merged.max_col))
        ws.unmerge_cells(str(merged))
    dimensions = {r: copy.copy(value) for r, value in ws.row_dimensions.items() if r >= row}
    for r in dimensions:
        del ws.row_dimensions[r]
    ws.insert_rows(row, count)
    for r, value in dimensions.items():
        value.index = r + count
        ws.row_dimensions[r + count] = value
    for first, last, left, right in merges:
        ws.merge_cells(start_row=first + count, end_row=last + count,
                       start_column=left, end_column=right)


def task_object(dialog, registry):
    char = registry['characters'][dialog['target']]
    title = dialog.get('task', {}).get('title') or char.get('talk_title')
    location = char.get('location') or registry.get('location')
    require(bool(title), f'Missing talk_title for {dialog["target"]}')
    require(bool(location), f'Missing location for {dialog["target"]}')
    fields = {'type': 'action', 'icon': char['classname'], 'action': dialog['action_identifier'],
              'amount': 1, 'title': title,
              'hint': dialog.get('task', {}).get('hint') or
                      f'{title}. Для этого просто кликни на персонажа. Локация: {location}.',
              'go_to_location': json.dumps([{'classname': char['classname']}], ensure_ascii=False,
                                          separators=(',', ':')),
              'identifier': dialog.get('task_identifier')}
    return fields


def insert_task(ws, row, output, index, fields, template):
    count = len(fields) + 4
    insert_rows(ws, row, count)
    for offset in range(count):
        styles = template[min(offset, 3)]
        for col, style in enumerate(styles, 1):
            ws.cell(row + offset, col)._style = copy.copy(style)
        ws.row_dimensions[row + offset].height = 21
    ws.cell(row, 2, f'Таск {index + 1}')
    ws.cell(row, 3, 'Диалог')
    for col, value in enumerate(['ml', 'string', 'string', 'object'], 1):
        ws.cell(row + 1, col, value)
    ws.cell(row + 2, 2, 'input')
    ws.cell(row + 2, 3, 'output')
    ws.cell(row + 2, 4, f'tasks.{index}')
    for offset, (key, value) in enumerate(fields.items(), 3):
        ws.cell(row + offset, 4, key)
        ws.cell(row + offset, 5, value)
        ws.cell(row + offset, 5).alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        if key in {'hint', 'go_to_location'}:
            ws.row_dimensions[row + offset].height = 48
    ws.cell(row + 3, 2, output)
    ws.cell(row + 3, 3, output)
    for offset in (1, 2):
        ws.merge_cells(start_row=row + offset, end_row=row + offset, start_column=4, end_column=5)


def update_quest_sheet(ws, dialogs, registry, quest_helpers=None):
    quests = scan_quests(ws)
    mapping = helper_views(registry)
    quest_helpers = quest_helpers or {}
    insertions, added, existing = [], 0, 0
    for name, quest in quests.items():
        row = quest['row']
        helper_col = next((c for c in range(4, ws.max_column + 1)
                           if ws.cell(row + 1, c).value == 'helper'), None)
        require(helper_col is not None, f'Missing helper column: {name}')
        old = ws.cell(row + 2, helper_col).value
        alias = quest_helpers.get(name, old)
        require(alias in mapping, f'Unknown helper for {name}: {alias!r}')
        ws.cell(row + 2, helper_col, mapping[alias])
        planned = [d for d in dialogs if d['quest'] == name]
        tasks = quest['tasks']
        require(bool(tasks), f'{name} has no task blocks; finish Stage 6 first')
        template = [[copy.copy(ws.cell(tasks[0]['row'] - 1 + off, col)._style)
                     for col in range(1, ws.max_column + 1)] for off in range(4)]
        missing_start = False
        for dialog in planned:
            matches = [(i, t) for i, t in enumerate(tasks)
                       if t['fields'].get('action') == dialog['action_identifier']]
            require(len(matches) <= 1, f'Duplicate dialog task: {dialog["action_identifier"]}')
            if matches:
                index, task = matches[0]
                expected = task_object(dialog, registry)
                for field in ('type', 'icon', 'action', 'amount'):
                    require(task['fields'].get(field) == expected[field],
                            f'Conflicting existing dialog task field {field}: {name}')
                require(json.loads(task['fields'].get('go_to_location', 'null')) ==
                        json.loads(expected['go_to_location']), f'Wrong dialog target: {name}')
                require(index == (0 if dialog['position'] == 'start' else len(tasks) - 1),
                        f'Dialog task is in the wrong position: {name}')
                if dialog.get('task_identifier'):
                    require(task['fields'].get('identifier') == dialog['task_identifier'],
                            f'Assigned dialog task ID differs from plan: {name}')
                existing += 1
            else:
                position = tasks[0]['row'] - 1 if dialog['position'] == 'start' else tasks[-1]['end'] + 1
                index = 0 if dialog['position'] == 'start' else len(tasks)
                if dialog['position'] == 'start':
                    missing_start = True
                insertions.append([position, quest['output'], index, dialog, template])
                added += 1
        if missing_start:
            for index, task in enumerate(tasks, 1):
                ws.cell(task['row'] + 1, 4, f'tasks.{index}')
                ws.cell(task['row'] - 1, 2, f'Таск {index + 1}')
            for item in insertions:
                if item[1] == quest['output'] and item[3]['position'] == 'end':
                    item[2] += 1
    if insertions:
        # Inserting rows cannot safely repair arbitrary Excel formulas referencing moved cells.
        for row in ws:
            for cell in row:
                require(cell.data_type != 'f' or not re.search(r'\$?[A-Z]+\$?\d+', cell.value),
                        f'Address formula needs manual reconciliation before row insertion: {cell.coordinate}')
    for row, output, index, dialog, template in sorted(insertions, key=lambda item: item[0], reverse=True):
        insert_task(ws, row, output, index, task_object(dialog, registry), template)
    return {'added_tasks': added, 'existing_tasks': existing}


def block_end(ws, row):
    end = row + 2
    while end <= ws.max_row and any(ws.cell(end, c).value is not None for c in range(1, ws.max_column + 1)):
        end += 1
    return end


def parse_array(ws, row, column):
    values = []
    for r in range(row + 2, block_end(ws, row)):
        value = ws.cell(r, column).value
        if value is not None:
            try:
                value = json.loads(value) if isinstance(value, str) else value
            except ValueError:
                pass
            values.append(value)
    return values


def write_array(ws, row, col, values):
    end = block_end(ws, row)
    if len(values) > end - row - 2:
        insert_rows(ws, end, len(values) - (end - row - 2))
    for r in range(row + 2, block_end(ws, row)):
        ws.cell(r, col).value = None
    for offset, value in enumerate(values, 2):
        ws.cell(row + offset, col, json.dumps(value, ensure_ascii=False, separators=(',', ':')))


def field_column(ws, row, key, typ):
    for col in range(4, ws.max_column + 1):
        if ws.cell(row + 1, col).value == key:
            require(ws.cell(row, col).value == typ, f'Wrong type for {key}')
            return col
    col = max(c for c in range(1, ws.max_column + 1)
              if ws.cell(row, c).value is not None or ws.cell(row + 1, c).value is not None) + 1
    for off in range(3):
        ws.cell(row + off, col)._style = copy.copy(ws.cell(row + off, 4)._style)
    ws.cell(row, col, typ)
    ws.cell(row + 1, col, key)
    return col


def bind_characters(ws, dialogs, registry, prototypes, snapshots=None):
    required = {}
    for dialog in dialogs:
        required.setdefault(registry['characters'][dialog['target']]['classname'], []).append(dialog['action_identifier'])
    found, outputs = set(), {}
    for row in range(ws.max_row - 2, 0, -1):
        if ws.cell(row, 1).value != 'ml':
            continue
        headers = {ws.cell(row + 1, c).value: c for c in range(4, ws.max_column + 1)
                   if ws.cell(row + 1, c).value is not None}
        col = headers.get('classname')
        classname = ws.cell(row + 2, col).value if col else None
        if classname not in required:
            continue
        found.add(classname)
        output = proto_path(prototypes, ws.cell(row + 2, 3).value)
        content = output.read_bytes() if output.exists() else None
        if snapshots is not None:
            snapshots[output] = content
        base = json.loads(content.decode('utf-8-sig')) if content is not None else read_json(
            proto_path(prototypes, ws.cell(row + 2, 2).value))
        char = next(c for c in registry['characters'].values() if c['classname'] == classname)
        root_col, legacy_col = headers.get('behaviour'), headers.get('behaviour.0.actions')
        behaviour = parse_array(ws, row, root_col) if root_col else copy.deepcopy(base.get('behaviour', []))
        if legacy_col and not root_col:
            require(bool(behaviour) and behaviour[0].get('type') == 'stuff_action',
                    f'Cannot safely reconcile behaviour.0.actions for {classname}')
            actions = list(dict.fromkeys([*parse_array(ws, row, legacy_col), *required[classname]]))
            write_array(ws, row, legacy_col, actions)
            behaviour[0]['actions'] = actions
        else:
            handler = next((b for b in behaviour if b.get('type') == 'stuff_action'), None)
            if handler is None:
                handler = {'type': 'stuff_action', 'actions': []}
                behaviour.append(handler)
            handler['actions'] = list(dict.fromkeys([*handler.get('actions', []), *required[classname]]))
            root_col = field_column(ws, row, 'behaviour', 'array')
            write_array(ws, row, root_col, behaviour)
        view_col = field_column(ws, row, 'view_classname', 'string')
        ws.cell(row + 2, view_col, char['view_class'])
        for field in ('classname', 'title', 'id'):
            if field in headers:
                value = ws.cell(row + 2, headers[field]).value
                if field == 'id' and output.exists():
                    require(value == base.get('id'), f'Character ID mismatch: {classname}')
                base[field] = value
        base['behaviour'] = behaviour
        base['view_classname'] = char['view_class']
        outputs[output] = base
    require(found == set(required), f'Missing characters in actions.xlsx: {sorted(set(required) - found)}')
    if ws.cell(1, 2).value == 'ПЕРСОНАЖИ БЕЗ ЭКШЕНОВ':
        ws.cell(1, 2, 'ПЕРСОНАЖИ')
    return outputs


def occupied_ids(files):
    """Read IDs conservatively, including uncached simple ID formulas in manual workbooks."""
    numbers, tasks = set(), set()
    for path in files:
        wb = load_workbook(path, data_only=False)
        try:
            ws = wb.active

            def number(cell, chain=()):
                if type(cell.value) is int:
                    return cell.value
                if cell.data_type == 'f':
                    match = re.fullmatch(r'=([A-Z]+[1-9]\d*)\+([0-9]+)', cell.value)
                    if match:
                        require(cell.coordinate not in chain, f'Cyclic ID formula: {path}:{cell.coordinate}')
                        base = number(ws[match[1]], (*chain, cell.coordinate))
                        return base + int(match[2]) if base is not None else None
                return None

            for row in ws:
                for cell in row:
                    value = number(cell)
                    if value is not None:
                        numbers.add(value)
                    if isinstance(cell.value, str) and re.fullmatch(r'e[1-9]\d*', cell.value):
                        tasks.add(cell.value)
        finally:
            wb.close()
    return numbers, tasks


def load_builder():
    path = PROJECT / '.agents/skills/quest-dialog-proto/scripts/build_dialog.py'
    spec = importlib.util.spec_from_file_location('quest_dialog_builder', path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def load_plan(path, campaign_id, pack_id, quests):
    plan = read_json(path)
    require(plan.get('version') == 1 and plan.get('campaign_id') == campaign_id and plan.get('pack_id') == pack_id,
            'Dialog plan version/campaign/pack does not match command')
    dialogs = plan.get('dialogs')
    require(isinstance(dialogs, list) and bool(dialogs), 'Dialog plan is empty')
    first = next(iter(quests))
    expected = {(q, 'end') for q in quests} | {(first, 'start')}
    actual = [(d.get('quest'), d.get('position')) for d in dialogs]
    require(len(actual) == len(set(actual)) and set(actual) == expected,
            'Pack needs one start dialog for its first quest and one end dialog per quest')
    return plan


def allocate(start, used, label):
    require(type(start) is int and start > 0, f'Provide {label} for new prototypes')
    while start in used:
        start += 1
    used.add(start)
    return start


def build_dialogs(plan, registry, prototypes, used, proto_start=None, snapshots=None):
    builder, outputs, seen_ids, seen_names = load_builder(), {}, set(), set()
    next_id = proto_start
    for index, entry in enumerate(plan['dialogs'], 1):
        target = entry.get('target')
        require(target in registry['characters'], f'Unknown dialog target: {target}')
        prefix = f'{plan["campaign_id"]}_{plan["pack_id"]}'
        entry.setdefault('file_name', f'{prefix}_Dialog_{index}.proto.js')
        entry.setdefault('action', {}).setdefault('file_name', f'{prefix}_{target}_Dialog_{index}.proto.js')
        for kind, spec, sub in [('UI', entry, 'user_interface/dialog'),
                                ('action', entry['action'], 'quest_action')]:
            name = spec['file_name']
            require(name.endswith('.proto.js'), f'Bad {kind} filename')
            stem = identifier(name.removesuffix('.proto.js'))
            require(stem.startswith(plan['campaign_id'] + '_'), f'Foreign {kind} filename: {name}')
            require(stem not in seen_names, f'Duplicate identifier: {stem}')
            seen_names.add(stem)
            path = proto_path(prototypes, f'/{sub}/{plan["campaign_id"]}/{name}')
            content = path.read_bytes() if path.exists() else None
            if snapshots is not None:
                snapshots[path] = content
            existing = json.loads(content.decode('utf-8-sig')) if content is not None else None
            assigned = spec.get('id')
            if existing:
                require(existing.get('identifier') == stem, f'Existing file has another identifier: {path}')
                require(assigned is None or assigned == existing.get('id'), f'Cannot reassign existing ID: {path}')
                assigned = existing.get('id')
            elif assigned is None:
                assigned = allocate(next_id, used, '--proto-start')
                next_id = assigned + 1
            else:
                require(type(assigned) is int and assigned > 0 and assigned not in used,
                        f'New proto ID already used or invalid: {assigned}')
                used.add(assigned)
            require(type(assigned) is int and assigned > 0 and assigned not in seen_ids,
                    f'Duplicate/invalid dialog proto ID: {assigned}')
            seen_ids.add(assigned)
            spec['id'] = assigned
        characters = []
        for item in entry.get('characters', []):
            key = item.get('key')
            require(key in registry['characters'], f'Unknown dialog speaker: {key}')
            char = registry['characters'][key]
            characters.append({'key': key, 'name': char['name'], 'side': item.get('side'),
                               'emotions': {'базовая': char['view_class'], **char.get('emotions', {})},
                               'asset_facing': char.get('asset_facing'),
                               'initial_emotion': item.get('initial_emotion', 'базовая')})
        conditions = f'active_quest={entry["quest"]}'
        require(entry['action'].get('conditions', conditions) == conditions,
                'Pack dialog actions use only active_quest; no task-completion restrictions')
        entry['action']['conditions'] = conditions
        brief = {**entry, 'characters': characters, 'expected_line_count': len(entry.get('lines', []))}
        ui_name, ui = builder.build(brief)
        action_name, action = builder.build_action(brief, ui['identifier'])
        for sub, name, obj in [('user_interface/dialog', ui_name, ui), ('quest_action', action_name, action)]:
            path = proto_path(prototypes, f'/{sub}/{plan["campaign_id"]}/{name}')
            require(not path.exists() or read_json(path) == obj,
                    f'Existing proto differs from plan; reconcile manual edits first: {path}')
            outputs[path] = obj
        entry['action_identifier'] = action['identifier']
    candidate = max([*seen_ids, (next_id or 1) - 1]) + 1
    while candidate in used:
        candidate += 1
    return outputs, candidate


@dataclass
class Change:
    path: Path
    data: bytes
    original: bytes | None


def commit_changes(changes, backup):
    changed = [c for c in changes if c.data != c.original]
    require(len({c.path.resolve() for c in changed}) == len(changed), 'Duplicate transaction target')
    for item in changed:
        require((item.path.read_bytes() if item.path.exists() else None) == item.original,
                f'File changed during preparation: {item.path}')
        if item.original is not None:
            proto_ids.check_replace_access(item.path)
    completed = []
    backup.mkdir(parents=True, exist_ok=True)
    try:
        for index, item in enumerate(changed):
            if item.original is not None:
                (backup / f'{index:03}_{item.path.name}').write_bytes(item.original)
            item.path.parent.mkdir(parents=True, exist_ok=True)
            temporary = item.path.with_name(item.path.name + '.codex-' + uuid.uuid4().hex)
            try:
                temporary.write_bytes(item.data)
                os.replace(temporary, item.path)
            finally:
                temporary.unlink(missing_ok=True)
            completed.append(item)
        for item in changed:
            require(item.path.read_bytes() == item.data, f'Write verification failed: {item.path}')
    except BaseException:
        for item in reversed(completed):
            if item.original is None:
                item.path.unlink(missing_ok=True)
            else:
                item.path.write_bytes(item.original)
        raise
    return len(changed)


def execute(campaign_id, pack_id, *, campaigns_dir, game_root, tables_dir=None,
            plan_path=None, characters_path=None, proto_start=None, task_start=None,
            apply=False, output_dir=None, node='node', converter=None):
    identifier(campaign_id)
    identifier(pack_id)
    campaign = Path(campaigns_dir) / campaign_id
    pack = campaign / pack_id
    game_root = Path(game_root).resolve()
    prototypes = game_root / 'server/prototypes'
    working = game_root / 'data_ingenering/GD/Quests_Story' / campaign_id
    tables = Path(tables_dir) if tables_dir else working if working.exists() else pack / 'review/proto'
    plan_path = Path(plan_path) if plan_path else pack / PLAN
    registry_path = Path(characters_path) if characters_path else campaign / REGISTRY
    registry_bytes = registry_path.read_bytes()
    registry = load_registry(registry_path)
    sources = [tables / f'{campaign_id}_quests.xlsx', tables / f'{campaign_id}_actions.xlsx']
    original = {p: p.read_bytes() for p in [*sources, plan_path]}
    try:
        proto_ids.read_assigned_ids(sources)
    except ValueError as exc:
        raise ValueError(f'Assign base proto/task IDs with proto_ids.py before Stage 6.1: {exc}') from exc
    catalogs = sorted(p for p in tables.glob('*.xlsx') if not p.name.startswith('~$'))
    catalog_bytes = {p: p.read_bytes() for p in catalogs}
    used_protos, used_tasks = occupied_ids(catalogs)
    books = [load_workbook(p, data_only=False) for p in sources]
    try:
        require(all(wb.sheetnames == ['conf'] for wb in books), 'Expected exactly one conf per workbook')
        snapshots = [sheet_snapshot(wb.active) for wb in books]
        quests = scan_quests(books[0].active)
        plan = load_plan(plan_path, campaign_id, pack_id, quests)
        existing_dialogs = {t['fields'].get('action'): t['fields'].get('identifier')
                            for q in quests.values() for t in q['tasks'] if t['fields'].get('action')}
        for entry in plan['dialogs']:
            action = entry.get('action', {}).get('file_name', '').removesuffix('.proto.js')
            task_id = entry.get('task_identifier')
            if task_id:
                require(re.fullmatch(r'e[1-9]\d*', task_id) is not None, f'Invalid task identifier: {task_id}')
                require(task_id not in used_tasks or existing_dialogs.get(action) == task_id,
                        f'New dialog task ID already used: {task_id}')
        proto_snapshots = {}
        products, next_proto = build_dialogs(plan, registry, prototypes, used_protos, proto_start,
                                             snapshots=proto_snapshots)
        stats = update_quest_sheet(books[0].active, plan['dialogs'], registry, plan.get('quest_helpers'))
        actors = bind_characters(books[1].active, plan['dialogs'], registry, prototypes,
                                 snapshots=proto_snapshots)
        if output_dir is None:
            (PROJECT / 'output').mkdir(exist_ok=True)
            output = Path(tempfile.mkdtemp(prefix=f'dialogs_{campaign_id}_{pack_id}_', dir=PROJECT / 'output'))
        else:
            output = Path(output_dir)
            require(not output.exists() or not any(output.iterdir()), 'Use an empty output directory')
            output.mkdir(parents=True, exist_ok=True)
        staged = output / 'tables'
        staged.mkdir()
        for source, wb, snapshot in zip(sources, books, snapshots):
            if sheet_snapshot(wb.active) == snapshot:
                (staged / source.name).write_bytes(original[source])
            else:
                wb.save(staged / source.name)
    finally:
        for wb in books:
            wb.close()
    missing_task_ids = [d for d in plan['dialogs'] if not d.get('task_identifier')]
    if task_start is None:
        # A complete existing pack can be checked again without providing new ranges.
        existing = scan_saved_tasks(staged / sources[0].name)
        require(all(existing[d['action_identifier']].get('identifier') for d in missing_task_ids),
                'Provide --task-start for new dialog tasks')
        task_start = f'e{max([int(value[1:]) for value in used_tasks] or [0]) + 1}'
    require(re.fullmatch(r'e[1-9]\d*', task_start) is not None, 'Invalid --task-start')
    ids = proto_ids.assign(staged, next_proto, task_start, backup_dir=output / 'id_assignment',
                           reserved_proto_ids=used_protos, reserved_task_ids=used_tasks)
    saved_tasks = scan_saved_tasks(staged / sources[0].name)
    for dialog in plan['dialogs']:
        dialog['task_identifier'] = saved_tasks[dialog['action_identifier']]['identifier']
        dialog.pop('action_identifier', None)
    checked = validate_proto_exports.run([staged / p.name for p in sources],
        converter or game_root / 'data_ingenering/tool/NodeTest/radmirxan/xlsx-to-json.js',
        prototypes, output / 'converted', node=node, assigned_ids=True)
    changes = [Change(p, (staged / p.name).read_bytes(), original[p]) for p in sources]
    changes.append(Change(plan_path, json_bytes(plan), original[plan_path]))
    for path, obj in {**products, **actors}.items():
        changes.append(Change(path, json_bytes(obj), proto_snapshots[path]))
    for quest, spec in quests.items():
        destination = proto_path(prototypes, spec['output'])
        converted = read_json(proto_path(output / 'converted', spec['output']))
        value = read_json(destination) if destination.exists() else copy.deepcopy(converted)
        require(value.get('id') == converted.get('id'), f'Quest ID differs from workbook: {quest}')
        value['tasks'], value['helper'] = converted['tasks'], converted['helper']
        changes.append(Change(destination, json_bytes(value), destination.read_bytes() if destination.exists() else None))
    # Persist reviewable candidate files for both modes, without touching the actual prototypes in preview mode.
    for item in changes:
        relative = item.path.relative_to(prototypes) if item.path.is_relative_to(prototypes) else Path('inputs') / item.path.name
        candidate = output / 'candidate' / relative
        candidate.parent.mkdir(parents=True, exist_ok=True)
        candidate.write_bytes(item.data)
    for path, content in catalog_bytes.items():
        require(path.read_bytes() == content, f'Table changed during preparation: {path}')
    require(registry_path.read_bytes() == registry_bytes, 'Character registry changed during preparation')
    count = commit_changes(changes, output / 'original') if apply else 0
    return {**stats, 'dialogs': len(plan['dialogs']), 'replicas': sum(len(d['lines']) for d in plan['dialogs']),
            'mode': 'applied' if apply else 'preview', 'changed_files': count,
            'candidate': str(output), 'converter': checked['status'],
            'next_proto_id': next_proto, 'next_task_id': ids['next_task_id']}


def scan_saved_tasks(path):
    wb = load_workbook(path, data_only=False)
    try:
        return {t['fields']['action']: t['fields'] for q in scan_quests(wb.active).values()
                for t in q['tasks'] if t['fields'].get('action')}
    finally:
        wb.close()
