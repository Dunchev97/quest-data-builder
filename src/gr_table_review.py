"""Check or repair GR text using the actual reviewed drop table, preserving mechanics."""
from __future__ import annotations

import argparse
from datetime import datetime
import json
from pathlib import Path
import re
import shutil
from time import time_ns

from openpyxl import load_workbook

try:
    from .gr_sources import GARBAGE_GR_TEMPLATES, garbage_hint, location_hint
    from .proto_ids import check_replace_access
    from .validate_proto_exports import expected_updates
except ImportError:
    from gr_sources import GARBAGE_GR_TEMPLATES, garbage_hint, location_hint
    from proto_ids import check_replace_access
    from validate_proto_exports import expected_updates

PROJECT = Path(__file__).resolve().parents[1]


def source_specs(campaign_id, resource_path, index):
    specs = {}
    for output, fields, _ in expected_updates(resource_path):
        if not output.startswith("/global_reward/"):
            continue
        name = fields.get("rand_reward.asset")
        if not isinstance(name, str) or not re.fullmatch(re.escape(campaign_id) + r"_GR_\d+", name):
            continue
        actions = fields.get("actions", [])
        garbage_actions = [a for a in actions if a in {"clean_garbage", "clean_garbage_in_guest"}]
        if not garbage_actions:
            continue
        assets = fields.get("assets", [])
        if len(garbage_actions) != 1:
            raise ValueError(f"{name}: ambiguous home/guest actions; review manually")
        garbage = [index["quest_ready_garbage"].get(source) for source in assets]
        if any(not item for item in garbage):
            raise ValueError(f"{name}: unknown garbage sources {assets!r}")
        tags = set(fields.get("location_tags", []))
        if not assets and not tags:
            raise ValueError(f"{name}: neither concrete garbage assets nor location_tags are defined")
        locations = [loc["title"] for loc in index["quest_ready_locations"].values()
                     if (not assets or set(assets).intersection(loc.get("garbage_assets", [])))
                     and (not tags or tags.intersection(loc.get("tags", [])))]
        mode = "guest" if garbage_actions[0].endswith("_in_guest") else "home"
        locations = ", ".join(dict.fromkeys(locations))
        if assets:
            hint = garbage_hint([item["title"] for item in garbage], mode, locations)
            template = "TT-014" if mode == "guest" else "TT-015"
        else:
            if mode != "home":
                raise ValueError(f"{name}: TT-013 is home only; location-wide guest GR needs a separate template")
            hint, template = location_hint(locations), "TT-013"
        spec = {"hint": hint, "template": template}
        if name in specs and specs[name] != spec:
            raise ValueError(f"{name}: conflicting drop definitions")
        specs[name] = spec
    return specs


def planned_edits(campaign_id, books, specs):
    changes = {}
    for path, wb in books:
        ws = wb["conf"]
        edits = {}
        for start in range(1, ws.max_row + 1):
            mode = ws.cell(start, 1).value
            if mode not in {"sl", "ml"}:
                continue
            end = start + 2
            while end <= ws.max_row and any(ws.cell(end, c).value not in (None, "") for c in range(1, ws.max_column + 1)):
                if ws.cell(end, 1).value in {"sl", "ml"}:
                    break
                end += 1
            keys = {ws.cell(start + 1, c).value: c for c in range(4, ws.max_column + 1)
                    if ws.cell(start, c).value == "string"}
            for row in range(start + 2, end):
                output = ws.cell(row, 3).value
                if isinstance(output, str) and output.startswith("/quest_item/") and "description" in keys:
                    name = ws.cell(row, keys["classname"]).value if "classname" in keys else Path(output).name.removesuffix(".proto.js")
                    if name in specs:
                        edits[ws.cell(row, keys["description"]).coordinate] = specs[name]["hint"]
            for col in range(4, ws.max_column + 1):
                if ws.cell(start, col).value != "object" or not re.fullmatch(r"tasks\.\d+", str(ws.cell(start + 1, col).value)):
                    continue
                cells = {ws.cell(row, col).value: ws.cell(row, col + 1) for row in range(start + 2, end)
                         if ws.cell(row, col).value}
                name = cells["classname"].value if "classname" in cells else None
                if name not in specs:
                    continue
                if "hint" not in cells:
                    raise ValueError(f"{path.name}: task for {name} has no hint field")
                edits[cells["hint"].coordinate] = specs[name]["hint"]
                label = ws.cell(start - 1, 3) if start > 1 else None
                if label and label.value in {"GR с мусора в локации дома", *[v[0] for v in GARBAGE_GR_TEMPLATES.values()]}:
                    edits[label.coordinate] = GARBAGE_GR_TEMPLATES[specs[name]["template"]][0]
        changes[path] = {coord: value for coord, value in edits.items() if ws[coord].value != value}
    return changes


def review(campaign_id, tables_dir, *, apply=False, index_path=PROJECT / "data/quest_ready_index.json"):
    folder = Path(tables_dir)
    paths = [folder / f"{campaign_id}_{kind}.xlsx" for kind in ("quests", "resources")]
    original = {p: p.read_bytes() for p in paths}
    books = [(p, load_workbook(p, data_only=False)) for p in paths]
    try:
        index = json.loads(Path(index_path).read_text(encoding="utf-8"))
        specs = source_specs(campaign_id, paths[1], index)
        changes = planned_edits(campaign_id, books, specs)
        changed = [p for p in paths if changes[p]]
        backup = None
        if apply and changed:
            for p in changed:
                check_replace_access(p)
            backup = PROJECT / "output" / ("gr_text_" + datetime.now().strftime("%Y%m%d_%H%M%S_%f"))
            backup.mkdir(parents=True)
            staged = {}
            for p, wb in books:
                if p not in changed:
                    continue
                (backup / p.name).write_bytes(original[p])
                ws = wb["conf"]
                for coord, value in changes[p].items():
                    ws[coord] = value
                stage = backup / ("updated_" + p.name)
                wb.save(stage)
                check = load_workbook(stage, data_only=False)
                before = load_workbook(backup / p.name, data_only=False)
                try:
                    for row in before["conf"]:
                        for cell in row:
                            actual = check["conf"][cell.coordinate]
                            if (actual.value != changes[p].get(cell.coordinate, cell.value)
                                    or actual._style != cell._style or actual.comment != cell.comment):
                                raise ValueError(f"Unexpected manual edit: {p.name}!{cell.coordinate}")
                    if (check["conf"].freeze_panes != before["conf"].freeze_panes
                            or str(check["conf"].merged_cells) != str(before["conf"].merged_cells)):
                        raise ValueError(f"Sheet formatting changed: {p.name}")
                finally:
                    check.close()
                    before.close()
                staged[p] = stage
            for p in paths:
                if p.read_bytes() != original[p]:
                    raise ValueError(f"Manual table changed during GR review: {p}")
            committed = []
            try:
                for p in changed:
                    check_replace_access(p)
                    local = p.with_name("." + p.name + f".{time_ns()}.tmp")
                    try:
                        shutil.copyfile(staged[p], local)
                        local.replace(p)
                    finally:
                        if local.exists():
                            local.unlink()
                    committed.append(p)
            except OSError:
                for p in committed:
                    p.write_bytes(original[p])
                raise
            remaining = review(campaign_id, folder, index_path=index_path)
            if any(remaining["changes"].values()):
                raise ValueError("GR text repair did not pass read-back validation")
        return {"sources": len(specs), "changes": {str(p): edits for p, edits in changes.items()},
                "backup": str(backup) if backup else None}
    finally:
        for _, wb in books:
            wb.close()


def validate(campaign_id, tables_dir):
    result = review(campaign_id, tables_dir)
    if any(result["changes"].values()):
        raise ValueError("GR source/template/text mismatch: " + json.dumps(result["changes"], ensure_ascii=False))
    return result


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("campaign_id")
    parser.add_argument("--tables-dir", type=Path, required=True)
    parser.add_argument("--apply", action="store_true", help="Repair only hint, description and task template labels")
    args = parser.parse_args(argv)
    try:
        result = review(args.campaign_id, args.tables_dir, apply=args.apply)
    except (ValueError, OSError) as exc:
        print(str(exc))
        return 1
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if args.apply or not any(result["changes"].values()) else 1


if __name__ == "__main__":
    raise SystemExit(main())
