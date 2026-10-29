"""Assign user-supplied free IDs directly to reviewed conf workbooks."""
from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import datetime
import os
from pathlib import Path
import re

from openpyxl import load_workbook


PROJECT = Path(__file__).resolve().parents[1]


def scan_sheet(ws):
    """Map IDs by output path and task IDs by (output path, tasks.N)."""
    protos, tasks, outputs = defaultdict(list), defaultdict(list), set()
    for start in range(1, ws.max_row + 1):
        mode = ws.cell(start, 1).value
        if mode not in {"sl", "ml"}:
            continue
        end = start + 2
        while end <= ws.max_row and any(ws.cell(end, c).value not in (None, "")
                                       for c in range(1, ws.max_column + 1)):
            end += 1
        inputs = [r for r in range(start + 2, end) if ws.cell(r, 2).value]
        for r in inputs:
            output = ws.cell(r, 3).value
            if not isinstance(output, str) or not output.startswith("/") or not output.endswith(".proto.js"):
                raise ValueError(f"{ws.title}!C{r}: invalid output")
            outputs.add(output)
            for c in range(4, ws.max_column + 1):
                typ, key = ws.cell(start, c).value, ws.cell(start + 1, c).value
                if key == "id":
                    if typ != "int":
                        raise ValueError(f"{ws.title}!{ws.cell(start,c).coordinate}: id type must be int")
                    protos[output].append(ws.cell(r, c))
                    if mode == "ml" and any(ws.cell(rr, c).value not in (None, "") for rr in range(r + 1, end)):
                        raise ValueError(f"{ws.title}!{ws.cell(r,c).coordinate}: id on a continuation row")
                if typ == "object" and re.fullmatch(r"tasks\.\d+", str(key)):
                    cells = [ws.cell(rr, c + 1) for rr in range(r, end) if ws.cell(rr, c).value == "identifier"]
                    if len(cells) != 1:
                        raise ValueError(f"{output}:{key}: expected one identifier field")
                    tasks[(output, key)].extend(cells)
                if typ == 'string' and re.fullmatch(r'tasks\.\d+\.identifier', str(key)):
                    tasks[(output, key.removesuffix('.identifier'))].append(ws.cell(r, c))
    return protos, tasks, outputs


def collect_slots(books):
    protos, tasks, outputs = defaultdict(list), defaultdict(list), set()
    for path, wb in books:
        if wb.sheetnames != ["conf"]:
            raise ValueError(f"{path}: expected exactly one conf sheet")
        p, t, o = scan_sheet(wb.active)
        outputs.update(o)
        for source, target in ((p, protos), (t, tasks)):
            for key, cells in source.items():
                target[key].extend((path, c) for c in cells)
    missing = outputs - protos.keys()
    if missing:
        raise ValueError(f"Missing prototype id fields: {sorted(missing)}")
    return protos, tasks


def assigned_values(slots, *, task=False, require_all=True):
    values, owners = {}, {}
    for key, cells in slots.items():
        present = set()
        for path, cell in cells:
            value = cell.value
            if value in (None, ""):
                if require_all:
                    raise ValueError(f"{path.name}!{cell.coordinate}: missing {'task identifier' if task else 'proto id'}")
                continue
            valid = (isinstance(value, str) and re.fullmatch(r"e[1-9]\d*", value)) if task else (
                type(value) is int and value > 0)
            if not valid:
                raise ValueError(f"{path.name}!{cell.coordinate}: invalid {'task identifier' if task else 'proto id'} {value!r}")
            present.add(value)
        if len(present) > 1:
            raise ValueError(f"Conflicting IDs for {key}: {present}")
        if present:
            value = next(iter(present))
            if value in owners and owners[value] != key:
                raise ValueError(f"Duplicate {'task identifier' if task else 'proto id'} {value}: {owners[value]} and {key}")
            owners[value], values[key] = key, value
    return values


def read_assigned_ids(files):
    books = [(Path(p), load_workbook(p, data_only=False)) for p in files]
    try:
        protos, tasks = collect_slots(books)
        return assigned_values(protos), assigned_values(tasks, task=True)
    finally:
        for _, wb in books:
            wb.close()


def assigned_id_errors(files):
    try:
        read_assigned_ids(files)
    except (ValueError, OSError) as exc:
        return [str(exc)]
    return []


def ordered_files(folder):
    def order(path):
        name = path.stem
        kinds = ("_technical_quests", "_quests", "_actions", "_resources", "_interactive_", "_funcollection_", "_accomplished")
        ranks = (1, 0, 2, 3, 4, 5, 6)
        return next(((rank, path.name.casefold()) for kind, rank in zip(kinds, ranks)
                     if kind in name), (7, path.name.casefold()))
    return sorted((p for p in Path(folder).glob("*.xlsx") if not p.name.startswith("~$")), key=order)


def check_replace_access(path):
    """Check Excel sharing locks before replacing any file in the batch."""
    if os.name != "nt":
        return
    import ctypes
    from ctypes import wintypes
    kernel = ctypes.WinDLL("kernel32", use_last_error=True)
    create = kernel.CreateFileW
    create.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID,
                       wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
    create.restype = wintypes.HANDLE
    handle = create(str(path.resolve()), 0x40000000 | 0x00010000, 7, None, 3, 0, None)
    if handle == wintypes.HANDLE(-1).value:
        raise OSError(ctypes.get_last_error(), f"File is locked or cannot be replaced: {path}")
    kernel.CloseHandle.argtypes = [wintypes.HANDLE]
    kernel.CloseHandle(handle)


def assign(folder, proto_start, task_start, *, backup_dir=None,
           reserved_proto_ids=(), reserved_task_ids=()):
    if type(proto_start) is not int or proto_start <= 0 or not re.fullmatch(r"e[1-9]\d*", task_start):
        raise ValueError("Provide a positive proto_start and task_start in e123 format")
    files = ordered_files(folder)
    if not files:
        raise ValueError(f"No XLSX files found in {folder}")
    books = [(p, load_workbook(p, data_only=False)) for p in files]
    try:
        protos, tasks = collect_slots(books)
        existing_protos = assigned_values(protos, require_all=False)
        existing_tasks = assigned_values(tasks, task=True, require_all=False)
        original = {p: p.read_bytes() for p in files}
        before_values = {p: {c.coordinate: c.value for row in wb.active for c in row}
                         for p, wb in books}
        changes = defaultdict(dict)
        next_proto, next_task = proto_start, int(task_start[1:])
        for slots, existing, is_task in ((protos, existing_protos, False), (tasks, existing_tasks, True)):
            used = set(existing.values()) | set(reserved_task_ids if is_task else reserved_proto_ids)
            for key, cells in slots.items():
                value = existing.get(key)
                if value is None:
                    if is_task:
                        while f"e{next_task}" in used:
                            next_task += 1
                        value, next_task = f"e{next_task}", next_task + 1
                    else:
                        while next_proto in used:
                            next_proto += 1
                        value, next_proto = next_proto, next_proto + 1
                    used.add(value)
                for path, cell in cells:
                    if cell.value != value:
                        changes[path][cell.coordinate] = value
                        cell.value = value
        proto_values, task_values = assigned_values(protos), assigned_values(tasks, task=True)
        while next_proto in set(proto_values.values()) | set(reserved_proto_ids):
            next_proto += 1
        while f"e{next_task}" in set(task_values.values()) | set(reserved_task_ids):
            next_task += 1
        backup = Path(backup_dir) if backup_dir else PROJECT / "output" / ("proto_ids_" + datetime.now().strftime("%Y%m%d_%H%M%S_%f"))
        staged = []
        if changes:
            backup.mkdir(parents=True, exist_ok=False)
            (backup / "original").mkdir()
            (backup / "staged").mkdir()
        for path, wb in books:
            if path not in changes:
                continue
            (backup / "original" / path.name).write_bytes(original[path])
            stage = backup / "staged" / path.name
            wb.save(stage)
            reopened = load_workbook(stage, data_only=False)
            try:
                ws = reopened.active
                for coord, old in before_values[path].items():
                    if ws[coord].value != changes[path].get(coord, old):
                        raise ValueError(f"Unexpected cell change at {path.name}!{coord}")
                    if ws[coord]._style != wb.active[coord]._style or ws[coord].comment != wb.active[coord].comment:
                        raise ValueError(f"Cell formatting/comment changed at {path.name}!{coord}")
                if ws.freeze_panes != wb.active.freeze_panes or set(map(str, ws.merged_cells.ranges)) != set(map(str, wb.active.merged_cells.ranges)):
                    raise ValueError(f"Sheet navigation/merges changed in {path.name}")
            finally:
                reopened.close()
            staged.append((path, stage))
        # Schema checks on the saved files, plus complete cross-file ID checks.
        try:
            from .converter_tables import validate
        except ImportError:
            from converter_tables import validate
        saved_files = [next((s for p, s in staged if p == path), path) for path in files]
        for path in saved_files:
            errors = validate(path, assigned_ids=True)
            if errors:
                raise ValueError(f"{path}: {errors}")
        read_assigned_ids(saved_files)
        for path, _ in staged:
            check_replace_access(path)
            if path.read_bytes() != original[path]:
                raise ValueError(f"Workbook changed during assignment: {path}")
        committed = []
        try:
            for path, stage in staged:
                stage.replace(path)
                committed.append(path)
        except OSError:
            for path in committed:
                path.write_bytes(original[path])
            raise
        read_assigned_ids(files)
        return {"files": len(files), "changed_files": len(changes), "prototypes": len(proto_values),
                "tasks": len(task_values), "proto_first": min(proto_values.values()),
                "proto_last": max(proto_values.values()), "next_proto_id": next_proto,
                "next_task_id": f"e{next_task}", "task_first": min(task_values.values(), key=lambda v: int(v[1:])) if task_values else None,
                "task_last": max(task_values.values(), key=lambda v: int(v[1:])) if task_values else None,
                "backup": str(backup) if changes else None}
    finally:
        for _, wb in books:
            wb.close()


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("folder", type=Path)
    parser.add_argument("--proto-start", type=int, required=True)
    parser.add_argument("--task-start", required=True)
    args = parser.parse_args(argv)
    try:
        result = assign(args.folder, args.proto_start, args.task_start)
    except (ValueError, OSError) as exc:
        print(str(exc))
        return 1
    print(result)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
