"""Apply the confirmed Adventure corrections to existing, manually edited XLSX.

Always stages, reopens and converts before changing the delivery. Backups and
converter output stay in output/. Existing IDs (including manual z IDs) survive.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path
import re
import shutil
import subprocess
import sys
from datetime import datetime

from openpyxl import load_workbook

PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT))
from src.converter_tables import validate_sheet
from src.proto_ids import scan_sheet, check_replace_access
from src.validate_proto_exports import expected_updates, table_paths, read_field, assert_subset
try:
    from .adventure_table_review import apply_windows, repair_quest, split_inputs, remove_box
except ImportError:
    from adventure_table_review import apply_windows, repair_quest, split_inputs, remove_box


def window_classes(path):
    wb = load_workbook(path, data_only=False)
    try:
        result = {}
        for row in wb.active:
            for cell in row:
                if cell.value in ('quest_confirm_window', 'quest_complete_window', 'quest_progress_window',
                                  'quest_rule_window', 'quest_window'):
                    value = wb.active.cell(cell.row, cell.column + 1).value
                    if isinstance(value, str) and value.startswith('{'):
                        value = {k: v for k, v in json.loads(value).items()
                                 if k in ('view_classname', 'quest_task_item')}
                    result[cell.value] = value
        if len(result) != 5:
            raise ValueError('Reference quest_group must expose all five windows')
        return result
    finally:
        wb.close()


def scalar_value(path, key):
    return expected_updates(path)[0][1][key]


def convert_pack(files, converter, prototypes, sandbox, node):
    """Check strict structure, then all authored fields in real converter output."""
    sandbox.mkdir(parents=True)
    produced = set()
    for path in files:
        wb = load_workbook(path, data_only=False)
        try:
            if wb.sheetnames != ['conf'] or wb.active.freeze_panes:
                raise ValueError(f'{path}: expected a single conf without freeze panes')
            errors = validate_sheet(wb.active, assigned_ids=True)
            if errors:
                raise ValueError(f'{path}: {errors}')
        finally:
            wb.close()
        for row, donor, output in table_paths(path):
            for value in (donor, output):
                if not value.startswith('/') or '..' in Path(value).parts:
                    raise ValueError(f'Unsafe path at {path}:{row}')
            if donor not in produced:
                source = prototypes / donor.lstrip('/')
                target = sandbox / donor.lstrip('/')
                target.parent.mkdir(parents=True, exist_ok=True)
                shutil.copyfile(source, target)
            produced.add(output)
    config = sandbox.parent / (sandbox.name + '.run.json')
    config.write_text(json.dumps({'root': str(sandbox.resolve()), 'converter': str(converter.resolve()),
                                 'files': [str(p.resolve()) for p in files]}), encoding='utf-8')
    process = subprocess.run([str(node), str(PROJECT / 'tools/validate_proto_exports.cjs'), str(config.resolve())],
                             capture_output=True, text=True, encoding='utf-8')
    if process.returncode:
        raise ValueError(process.stderr or process.stdout)
    results = [json.loads(line) for line in process.stdout.splitlines() if line.strip()]
    fields_checked = 0
    for path, result in zip(files, results, strict=True):
        if len(result['writes']) != len(table_paths(path)):
            raise ValueError(f'{path}: converter silently omitted a prototype update')
        for output, fields, replacements in expected_updates(path):
            actual = json.loads((sandbox / output.lstrip('/')).read_text(encoding='utf-8'))
            for key, value in fields.items():
                assert_subset(read_field(actual, key), value, f'{output}:{key}')
                fields_checked += 1
            encoded = json.dumps(actual, ensure_ascii=False)
            for old, new in replacements:
                if old != new and old not in new and old in encoded:
                    raise ValueError(f'{output}: donor reference survived: {old}')
        wb = load_workbook(path)
        try:
            _, task_slots, _ = scan_sheet(wb.active)
            for (output, task), cells in task_slots.items():
                actual = json.loads((sandbox / output.lstrip('/')).read_text(encoding='utf-8'))
                if read_field(actual, task + '.identifier') != cells[0].value:
                    raise ValueError(f'{output}:{task}: task ID lost during conversion')
            for output in {o for _, _, o in table_paths(path) if o.startswith('/quest/')}:
                actual = json.loads((sandbox / output.lstrip('/')).read_text(encoding='utf-8'))
                for i, task in enumerate(actual.get('tasks', [])):
                    if (output, f'tasks.{i}') not in task_slots:
                        raise ValueError(f'{output}: unexpected inherited donor task {i}')
        finally:
            wb.close()
    return len(produced), fields_checked


def apply_review(root, event, reference_pack, task_start, work, converter, prototypes, node='node', *, apply=False):
    root, work, prototypes, converter = [Path(p).resolve() for p in (root, work, prototypes, converter)]
    if work == root or root in work.parents or work in root.parents:
        raise ValueError('Staging workspace must be separate from the reviewed Adventure directory')
    if work.exists():
        raise FileExistsError(work)
    folders = sorted(root.glob(f'{event}_*_pack_tables'), key=lambda p: int(p.name.split('_')[-3]))
    if not folders:
        raise ValueError('No Adventure pack tables found')
    windows = window_classes(root / f'{event}_{reference_pack}_pack_tables/quest_group.xlsx')
    original, existing_tasks, task_owners = {}, {}, {}
    for folder in folders:
        for path in folder.glob('*.xlsx'):
            if path.name.startswith('~$'):
                continue
            original[path] = path.read_bytes()
            if path.stem not in ('quest', 'trophy'):
                continue
            wb = load_workbook(path)
            try:
                for key, cells in scan_sheet(wb.active)[1].items():
                    for cell in cells:
                        value = cell.value
                        if value in (None, ''):
                            continue
                        if value in task_owners and task_owners[value] != key:
                            raise ValueError(f'Duplicate existing task identifier {value}')
                        if key in existing_tasks and existing_tasks[key] != value:
                            raise ValueError(f'Conflicting existing task IDs for {key}')
                        task_owners[value], existing_tasks[key] = key, value
            finally:
                wb.close()
    if not re.fullmatch(r'e[1-9]\d*', task_start):
        raise ValueError('task_start must be the explicitly supplied first free e identifier')
    next_number = int(task_start[1:])
    added = []
    def allocate():
        nonlocal next_number
        while f'e{next_number}' in task_owners:
            next_number += 1
        value = f'e{next_number}'
        next_number += 1
        added.append(value)
        return value

    work.mkdir(parents=True)
    changed, deleted = [], []
    total_tasks, total_protos, total_fields = 0, 0, 0
    all_after = {}
    for folder in folders:
        stage = work / 'staged' / folder.name
        stage.mkdir(parents=True)
        pet_title = scalar_value(folder / 'pet.xlsx', 'title')
        main_resource = scalar_value(folder / 'quest_item.xlsx', 'classname') if (folder / 'quest_item.xlsx').exists() else None
        for path in sorted(folder.glob('*.xlsx')):
            if path.name.startswith('~$'):
                continue
            if path.name == 'quest_item.xlsx':
                deleted.append(path)
                continue
            target = stage / path.name
            shutil.copyfile(path, target)
            if path.stem not in ('quest', 'trophy', 'mystery_box', 'quest_group'):
                continue
            wb = load_workbook(target, data_only=False)
            try:
                ws = wb.active
                before = [(c.coordinate, c.value) for row in ws for c in row if c.value is not None]
                merges = set(map(str, ws.merged_cells.ranges))
                if path.stem in ('quest', 'trophy'):
                    repair_quest(ws, pet_title, allocate, main_resource=main_resource)
                elif path.stem == 'mystery_box':
                    remove_box(ws, event + '_Box_1')
                elif path.stem == 'quest_group':
                    apply_windows(ws, windows)
                split_inputs(ws)
                after = [(c.coordinate, c.value) for row in ws for c in row if c.value is not None]
                if before != after or merges != set(map(str, ws.merged_cells.ranges)):
                    wb.save(target)
                    changed.append((path, target))
            finally:
                wb.close()
        files = sorted(stage.glob('*.xlsx'))
        for path in files:
            wb = load_workbook(path, data_only=False)
            try:
                for key, cells in scan_sheet(wb.active)[1].items():
                    values = {c.value for c in cells}
                    if len(values) != 1:
                        raise ValueError(f'{key}: inconsistent task IDs')
                    value = next(iter(values))
                    if value in (None, '') or (value not in task_owners and not re.fullmatch(r'e[1-9]\d*', str(value))):
                        raise ValueError(f'{key}: missing or invalid task ID')
                    if key in existing_tasks and existing_tasks[key] != value:
                        raise ValueError(f'{key}: a manual task identifier changed')
                    if value in all_after and all_after[value] != key:
                        raise ValueError(f'Duplicate task identifier {value}')
                    all_after[value] = key
                    total_tasks += 1
            finally:
                wb.close()
        count, fields = convert_pack(files, converter, prototypes, work / 'converted' / folder.name, node)
        total_protos += count
        total_fields += fields
        print(f'CHECKED {folder.name}: {len(files)} workbooks, {count} prototypes', flush=True)
    if apply:
        affected = [p for p, _ in changed] + deleted
        for path in affected:
            if path.parent not in folders or path.resolve().parent.parent != root:
                raise ValueError(f'Unexpected mutation path: {path}')
            check_replace_access(path)
            if path.read_bytes() != original[path]:
                raise ValueError(f'Workbook changed during review: {path}')
            backup = work / 'original' / path.relative_to(root)
            backup.parent.mkdir(parents=True, exist_ok=True)
            backup.write_bytes(original[path])
        committed = []
        try:
            for path, stage in changed:
                stage.replace(path)
                committed.append(path)
            for path in deleted:
                path.unlink()
                committed.append(path)
        except OSError:
            for path in committed:
                path.write_bytes(original[path])
            raise
    while f'e{next_number}' in all_after:
        next_number += 1
    return {'packs': len(folders), 'changed_files': len(changed), 'deleted_files': len(deleted),
            'new_task_ids': len(added), 'task_first': added[0] if added else None,
            'task_last': added[-1] if added else None, 'next_task_id': f'e{next_number}',
            'tasks_checked': total_tasks, 'prototypes_checked': total_protos,
            'fields_checked': total_fields, 'applied': apply, 'workspace': str(work)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('root', type=Path)
    parser.add_argument('--event', required=True)
    parser.add_argument('--reference-pack', type=int, default=2)
    parser.add_argument('--task-start', required=True)
    parser.add_argument('--work', type=Path)
    parser.add_argument('--converter', type=Path, required=True)
    parser.add_argument('--prototypes', type=Path, required=True)
    parser.add_argument('--node', default='node')
    parser.add_argument('--apply', action='store_true')
    args = parser.parse_args()
    work = args.work or PROJECT / 'output' / ('adventure_review_' + datetime.now().strftime('%Y%m%d_%H%M%S'))
    result = apply_review(args.root, args.event, args.reference_pack, args.task_start, work,
                          args.converter, args.prototypes, args.node, apply=args.apply)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
