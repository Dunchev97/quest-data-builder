"""Repairs shared by Adventure generation and migration of reviewed tables."""
from __future__ import annotations

from copy import copy
import json
import re

from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

DIRECTIVES = {'sl', 'ml'}
FAIRY_TALE_WINDOWS = {
    'quest_confirm_window': 'FairyTaleConfirmWindow',
    'quest_complete_window': 'FairyTaleQuestCompleteWindow',
    'quest_progress_window': 'FairyTaleQuestProgressWindow',
    'quest_rule_window': {'view_classname': 'FairyTaleQuestRuleWindow'},
    'quest_window': {
        'view_classname': 'FairyTaleQuestWindow',
        'quest_task_item': 'FairyTaleQuestTaskItem',
    },
}


def blocks(ws):
    for start in range(1, ws.max_row + 1):
        if ws.cell(start, 1).value not in DIRECTIVES:
            continue
        end = start + 2
        while end <= ws.max_row and any(c.value is not None for c in ws[end]):
            end += 1
        yield start, end


def scalar_field(ws, start, key, value=None):
    """Append a contiguous typed column; retain an existing manual value."""
    cols = [c.column for c in ws[start + 1] if c.value is not None]
    col = next((c for c in cols if ws.cell(start + 1, c).value == key), max(cols) + 1)
    for row, content in ((start, 'string'), (start + 1, key)):
        cell = ws.cell(row, col)
        if cell.value is None:
            cell._style = copy(ws.cell(row, max(4, col - 1))._style)
            cell.value = content
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    cell = ws.cell(start + 2, col)
    if cell.value in (None, '') and value is not None:
        cell._style = copy(ws.cell(start + 2, max(4, col - 1))._style)
        cell.value = value
    ws.column_dimensions[get_column_letter(col)].width = max(
        ws.column_dimensions[get_column_letter(col)].width or 13, len(key) + 3)
    return cell


def repair_quest(ws, pet_title, next_task=None, *, main_resource=None):
    """Expose every non-reward task and update rewards from the created quest."""
    for start, end in list(blocks(ws)):
        path = ws.cell(start + 2, 3).value
        if not isinstance(path, str) or not path.startswith('/quest/'):
            continue
        task_objects = [c for c in ws[start + 1]
                        if re.fullmatch(r'tasks\.\d+', str(c.value))]
        if task_objects:
            ws.cell(start + 2, 2).value = path
            continue
        if path.endswith('_rewards.proto.js'):
            continue
        if path.startswith('/quest/trophy/'):
            for col in range(4, ws.max_column + 1):
                if ws.cell(start, col).value == 'replace':
                    old, new = ws.cell(start + 2, col), ws.cell(start + 2, col + 1)
                    if isinstance(old.value, str) and old.value.startswith('TrophyPet'):
                        old.value = old.value.removeprefix('Trophy')
                        new.value = new.value.removeprefix('Trophy')
        identifier = scalar_field(ws, start, 'tasks.0.identifier')
        if identifier.value in (None, '') and next_task is not None:
            identifier.value = next_task()
        scalar_field(ws, start, 'tasks.0.title', pet_title)
        # The root has both a reward and remove_asset for this resource.
        if main_resource and path.endswith('_quest.proto.js'):
            for col in range(4, ws.max_column + 1):
                if ws.cell(start, col).value != 'replace':
                    continue
                if any(ws.cell(r, col).value == 'Adventure_GR_3'
                       for r in range(start + 2, end)):
                    break
                if ws.cell(start + 3, col).value is None:
                    ws.cell(start + 3, col).value = 'Adventure_GR_3'
                    ws.cell(start + 3, col + 1).value = main_resource
                    break


def apply_windows(ws, windows=None):
    windows = windows or FAIRY_TALE_WINDOWS
    for row in ws:
        for cell in row:
            if not isinstance(cell.value, str) or cell.value not in windows:
                continue
            value = windows[cell.value]
            target = ws.cell(cell.row, cell.column + 1)
            if isinstance(value, dict):
                current = json.loads(target.value or '{}')
                current.update(value)
                target.value = json.dumps(current, ensure_ascii=False, separators=(',', ':'))
            else:
                target.value = value


def _insert_rows(ws, index, count):
    merges = [copy(m) for m in ws.merged_cells.ranges]
    for merged in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged))
    dimensions = {r: copy(d) for r, d in ws.row_dimensions.items()}
    ws.row_dimensions.clear()
    ws.insert_rows(index, count)
    for row, dim in dimensions.items():
        new_row = row + count if row >= index else row
        dim.index = new_row
        ws.row_dimensions[new_row] = dim
    for merged in merges:
        if merged.min_row >= index:
            merged.shift(row_shift=count)
        elif merged.max_row >= index:
            merged.max_row += count
        ws.merge_cells(str(merged))


def split_inputs(ws):
    """ml and sl+replace need one prototype per block in the real converter."""
    plans = []
    for start, end in blocks(ws):
        if ws.cell(start, 1).value != 'ml' and not any(c.value == 'replace' for c in ws[start]):
            continue
        inputs = [r for r in range(start + 2, end) if ws.cell(r, 2).value]
        for row in inputs[1:]:
            plans.append((row, start))
    for row, start in sorted(plans, reverse=True):
        source = [[copy(c) for c in ws[r]] for r in (start, start + 1)]
        header_merges = [(m.min_row - start, m.min_col, m.max_col)
                         for m in ws.merged_cells.ranges if start <= m.min_row <= m.max_row <= start + 1]
        _insert_rows(ws, row, 3)
        for offset, cells in enumerate(source):
            for old in cells:
                target = ws.cell(row + 1 + offset, old.column)
                target.value = old.value
                target._style = copy(old._style)
                if old.comment:
                    target.comment = copy(old.comment)
        for offset, left, right in header_merges:
            ws.merge_cells(start_row=row + 1 + offset, end_row=row + 1 + offset,
                           start_column=left, end_column=right)


def remove_box(ws, classname):
    """Remove a reviewed box definition, preserving the other boxes' cells."""
    for row in range(ws.max_row, 1, -1):
        if ws.cell(row, 4).value == classname and ws.cell(row, 2).value:
            ws.delete_rows(row)
