import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from src import gr_table_review


class GrTableReviewTests(unittest.TestCase):
    def fixture(self, folder):
        quest = folder / "Event_2026_quests.xlsx"
        resources = folder / "Event_2026_resources.xlsx"
        rows = [
            [None, "Таск 1", "GR с мусора в локации дома", None, None, "=1+2"],
            ["ml", "string", "string", "object"],
            [None, "input", "output", "tasks.0"],
            [None, "/quest/New.proto.js", "/quest/New.proto.js", "type", "get_asset"],
            [None, None, None, "classname", "Event_2026_GR_1"],
            [None, None, None, "hint", "Убирай мусор в локации Котельная дома, чтобы найти."],
            [None, None, None, "identifier", "e777"],
            [None, None, None, "amount", 37],
        ]
        res_rows = [
            ["sl", "string", "string", "string", "string", "int"],
            [None, "input", "output", "classname", "description", "id"],
            [None, "/quest_item/Donor.proto.js", "/quest_item/Event_2026_GR_1.proto.js",
             "Event_2026_GR_1", "Убирай мусор дома", 123], [],
            ["sl", "string", "string", "array", "string", "array", "array"],
            [None, "input", "output", "actions", "rand_reward.asset", "assets", "location_tags"],
            [None, "/global_reward/Donor.proto.js", "/global_reward/Event_2026_GR_1.proto.js",
             "clean_garbage_in_guest", "Event_2026_GR_1", "RustyNail"],
        ]
        for path, data in ((quest, rows), (resources, res_rows)):
            wb = Workbook()
            ws = wb.active
            ws.title = "conf"
            for row in data:
                ws.append(row)
            ws.freeze_panes = "A11"  # Existing manual settings must survive text-only repair.
            ws["A1"].comment = Comment("Ручной комментарий", "GD")
            wb.save(path)
            wb.close()
        index = folder / "index.json"
        index.write_text(json.dumps({"quest_ready_garbage": {"RustyNail": {"title": "Ржавый гвоздь"}},
                                     "quest_ready_locations": {"65": {"title": "Котельная", "tags": ["boiler"],
                                                                            "garbage_assets": ["RustyNail"]}}},
                                    ensure_ascii=False), encoding="utf-8")
        return quest, resources, index

    def test_actual_drop_mode_wins_and_only_text_cells_change(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            quest, resources, index = self.fixture(folder)
            before = {p: p.read_bytes() for p in (quest, resources)}
            check = gr_table_review.review("Event_2026", folder, index_path=index)
            self.assertEqual(sum(map(len, check["changes"].values())), 3)
            self.assertEqual({p: p.read_bytes() for p in before}, before)
            result = gr_table_review.review("Event_2026", folder, apply=True, index_path=index)
            expected = "Убирай Ржавый гвоздь в гостях, чтобы найти. Место поиска: Котельная."
            for p, coord in ((quest, "E6"), (resources, "E3")):
                wb = load_workbook(p, data_only=False)
                self.assertEqual(wb["conf"][coord].value, expected)
                self.assertEqual(wb["conf"].freeze_panes, "A11")
                self.assertEqual(wb["conf"]["A1"].comment.text, "Ручной комментарий")
                wb.close()
                self.assertEqual((Path(result["backup"]) / p.name).read_bytes(), before[p])
            wb = load_workbook(quest, data_only=False)
            self.assertEqual(wb["conf"]["E7"].value, "e777")
            self.assertEqual(wb["conf"]["E8"].value, 37)
            self.assertEqual(wb["conf"]["F1"].value, "=1+2")
            wb.close()
            stable = {p: p.read_bytes() for p in before}
            self.assertFalse(any(gr_table_review.review("Event_2026", folder, apply=True, index_path=index)["changes"].values()))
            self.assertEqual({p: p.read_bytes() for p in before}, stable)

    def test_location_wide_drop_requires_tags_and_has_no_concrete_source_hint(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            _, resources, index = self.fixture(folder)
            wb = load_workbook(resources, data_only=False)
            ws = wb["conf"]
            ws["D7"] = "clean_garbage"
            ws["F7"] = None
            wb.save(resources)
            with self.assertRaisesRegex(ValueError, "neither concrete"):
                gr_table_review.review("Event_2026", folder, index_path=index)
            ws["G7"] = "boiler"
            wb.save(resources)
            wb.close()
            specs = gr_table_review.source_specs("Event_2026", resources, json.loads(index.read_text(encoding="utf-8")))
            self.assertEqual(specs["Event_2026_GR_1"], {"template": "TT-013",
                "hint": "Убирай мусор в локации Котельная дома, чтобы найти."})


if __name__ == "__main__":
    unittest.main()
