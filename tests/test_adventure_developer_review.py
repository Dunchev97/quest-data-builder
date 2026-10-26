from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from src.converter_tables import validate_sheet
from src.proto_ids import scan_sheet, assign
from src.validate_proto_exports import expected_updates, table_paths
from tools import adventure_table_review as review
from tools import build_adventure_pack_from_brief as build


class AdventureDeveloperReviewTests(unittest.TestCase):
    def test_adds_missing_root_restarter_tech_and_trophy_without_changing_existing_ids(self):
        wb = Workbook()
        ws = wb.active
        ws.title = 'conf'
        for name, identifier in [('Event_quest', 'z10101'), ('Event_restarter', None),
                                 ('Event_tech', None), ('TrophyPet17_Event', None)]:
            ws.append(['ml', 'string', 'string', 'int', 'string'])
            ws.append([None, 'input', 'output', 'id', 'tasks.0.identifier'])
            ws.append([None, '/quest/donor.proto.js', f'/quest/{name}.proto.js', 123, identifier])
            ws.append([])
        ids = iter(['e100317', 'e100318', 'e100319'])
        review.repair_quest(ws, 'Мышка', lambda: next(ids))
        tasks = scan_sheet(ws)[1]
        self.assertEqual([v[0].value for v in tasks.values()], ['z10101', 'e100317', 'e100318', 'e100319'])
        review.repair_quest(ws, 'Мышка', lambda: self.fail('Repeated repair allocated an ID'))
        self.assertEqual(validate_sheet(ws, assigned_ids=True), [])

    def test_reward_updates_read_created_quest_and_preserve_manual_text(self):
        wb = Workbook()
        ws = wb.active
        ws.title = 'conf'
        for i in range(2):
            ws.append(['ml', 'string', 'string', 'object', None])
            ws.append([None, 'input', 'output', f'tasks.{i}', None])
            ws.append([None, '/quest/Donor_rewards.proto.js', '/quest/Event_rewards.proto.js', 'title', 'Ручная правка'])
            ws.append([None, None, None, 'identifier', f'e{500+i}'])
            ws.append([])
        review.repair_quest(ws, 'Мышка')
        self.assertEqual(ws['B3'].value, '/quest/Event_rewards.proto.js')
        self.assertEqual(ws['B8'].value, '/quest/Event_rewards.proto.js')
        self.assertEqual(ws['E3'].value, 'Ручная правка')

    def test_split_retains_all_box_objects_and_replacements_after_reopen(self):
        for directive, typ in [('ml', 'object'), ('sl', 'replace')]:
            wb = Workbook()
            ws = wb.active
            ws.title = 'conf'
            ws.append([directive, 'string', 'string', typ, None])
            ws.append([None, 'input', 'output', 'rand_reward' if typ == 'object' else 'find', None if typ == 'object' else 'replace'])
            for i in range(3):
                ws.append([None, '/donor.proto.js', f'/new_{i}.proto.js', 'all' if typ == 'object' else 'Old', f'Value_{i}'])
            ws.merge_cells('D1:E1')
            if typ == 'object':
                ws.merge_cells('D2:E2')
            ws.row_dimensions[4].height = 50
            review.split_inputs(ws)
            self.assertEqual(validate_sheet(ws, assigned_ids=True), [])
            with tempfile.TemporaryDirectory() as tmp:
                p = Path(tmp) / 'boxes.xlsx'
                wb.save(p)
                reopened = load_workbook(p)
                self.assertEqual(validate_sheet(reopened.active, assigned_ids=True), [])
                self.assertEqual([donor for _, donor, _ in table_paths(p)], ['/donor.proto.js'] * 3)
                self.assertEqual(len(expected_updates(p)), 3)
                self.assertEqual(reopened.active.row_dimensions[7].height, 50)

    def test_windows_keep_manual_rule_text_and_extra_keys(self):
        wb = Workbook()
        ws = wb.active
        ws.append(['quest_rule_window', '{"view_classname":"Old","title":"Правила","description":"Ручной текст"}'])
        ws.append(['quest_window', '{"view_classname":"Old","quest_task_item":"OldItem","custom":5}'])
        review.apply_windows(ws)
        import json
        self.assertEqual(json.loads(ws['B1'].value)['description'], 'Ручной текст')
        self.assertEqual(json.loads(ws['B1'].value)['view_classname'], 'FairyTaleQuestRuleWindow')
        self.assertEqual(json.loads(ws['B2'].value)['custom'], 5)

    def test_generator_exposes_all_seventeen_tasks_and_trophy_pet_link(self):
        data = {
            'pack': {'new_prefix': 'Adventure_8_1', 'main_resource': 'Adventure_GR_8', 'event_title_ru': 'Сказка'},
            'event_folder': 'Adventure_8', 'quests_by_key': {}, 'objects': [],
            'objects_by_kind': {'pet': {'title': 'Мышка', 'new_classname': 'Pet17_Adventure_8_1'},
                                'trophy': {'title': 'Мышка', 'new_classname': 'TrophyPet17_Adventure_8_1'}},
            'task_ids': build.TaskIdAllocator('e100316'), 'proto_ids': build.ProtoIdAllocator(None),
            'rewards': [{'classname': 'Adventure_GR_8', 'amount': (i+1)*500,
                         'reward': 'asset=ExistingItem:1', 'title': f'Награда {i}', 'hint': 'Описание',
                         'congratulation': 'Спасибо'} for i in range(7)],
        }
        mapping = {build.OLD_PREFIX: 'Adventure_8_1', build.OLD_EVENT_FOLDER: 'Adventure_8',
                   build.OLD_MAIN_RESOURCE: 'Adventure_GR_8', 'Pet17_Sport_9': 'Pet17_Adventure_8_1'}
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            paths = [build.build_quest(data, mapping, root, []), build.build_trophy(data, mapping, root)]
            identifiers = []
            for path in paths:
                wb = load_workbook(path)
                self.assertEqual(validate_sheet(wb.active), [])
                self.assertIsNone(wb.active.freeze_panes)
                identifiers.extend(cells[0].value for cells in scan_sheet(wb.active)[1].values())
            self.assertEqual(len(identifiers), 17)
            self.assertEqual(len(set(identifiers)), 17)
            self.assertEqual(set(identifiers), {f'e{i}' for i in range(100317, 100334)})
            updates = expected_updates(paths[1])[0]
            self.assertEqual(updates[2], [('Pet17_Sport_9', 'Pet17_Adventure_8_1')])
            self.assertEqual(updates[1]['tasks.0.title'], 'Мышка')
            root_replacements = expected_updates(paths[0])[0][2]
            self.assertIn(('Adventure_GR_3', 'Adventure_GR_8'), root_replacements)
            for output, fields, _ in expected_updates(paths[0]):
                if any(k.startswith('tasks.') and isinstance(v, dict) for k, v in fields.items()):
                    self.assertTrue(output.endswith('Adventure_8_1_rewards.proto.js'))

    def test_scalar_task_ids_participate_in_assignment_and_uniqueness(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp) / 'tables'
            folder.mkdir()
            path = folder / 'quest.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.title = 'conf'
            ws.append(['sl', 'string', 'string', 'int', 'string', 'string'])
            ws.append([None, 'input', 'output', 'id', 'tasks.0.identifier', 'tasks.0.title'])
            ws.append([None, '/quest/old.proto.js', '/quest/new.proto.js', None, None, 'Мышка'])
            for row in ws:
                for cell in row:
                    cell.font = Font(name='Calibri', size=11)
            wb.save(path)
            result = assign(folder, 100, 'e100317', backup_dir=Path(tmp) / 'backup')
            self.assertEqual(result['tasks'], 1)
            self.assertEqual(result['next_task_id'], 'e100318')
            self.assertEqual(load_workbook(path).active['E3'].value, 'e100317')

    def test_human_proto_titles_are_not_converter_inputs(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'table.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.append([None, 'main - donor.proto.js'])
            ws.append(['sl', 'string', 'string', 'string'])
            ws.append([None, 'input', 'output', 'title'])
            ws.append([None, '/quest/donor.proto.js', '/quest/new.proto.js', 'Мышка'])
            wb.save(path)
            self.assertEqual(table_paths(path), [(4, '/quest/donor.proto.js', '/quest/new.proto.js')])


if __name__ == '__main__':
    unittest.main()
