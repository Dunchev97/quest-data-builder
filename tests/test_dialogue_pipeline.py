import copy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

from src import dialogue_pipeline as dialogs
from src import proto_ids, workflow_fast


def registry():
    return {'version': 1, 'location': 'Тёплые истории', 'characters': {
        'ded': {'name': 'Дед Домовед', 'classname': 'Event_Character_1_1',
                'view_class': 'Oseniny_Character_1', 'emotions': {'радость': 'Oseniny_Character_1_Happy'},
                'asset_facing': {'Oseniny_Character_1': 'left', 'Oseniny_Character_1_Happy': 'left'},
                'talk_title': 'Поговори с Дедом Домоведом'}}}


def fixture_quests():
    wb = Workbook()
    ws = wb.active
    ws.title = 'conf'
    output = '/quest/event/Event/Event_Story_1.proto.js'
    for row in [
        [None, 'Quest'], ['sl', 'string', 'string', 'string', 'int'],
        [None, 'input', 'output', 'helper', 'id'],
        [None, '/quest/donor.proto.js', output, 'Event_Character_1_1', 11], [],
    ]:
        ws.append(row)
    for index in range(2):
        for row in [[None, f'Таск {index + 1}', 'Manual task'],
                    ['ml', 'string', 'string', 'object'],
                    [None, 'input', 'output', f'tasks.{index}'],
                    [None, output, output, 'type', 'get_asset'],
                    [None, None, None, 'classname', f'Old_{index}'],
                    [None, None, None, 'amount', index + 3],
                    [None, None, None, 'identifier', f'e{21 + index}'], []]:
            ws.append(row)
        first = ws.max_row - 5
        ws.merge_cells(start_row=first, end_row=first, start_column=4, end_column=5)
        ws.merge_cells(start_row=first + 1, end_row=first + 1, start_column=4, end_column=5)
    ws['E11'].font = Font(bold=True, color='AA0000')
    ws['E11'].comment = Comment('Ручная правка', 'User')
    ws.freeze_panes = 'A5'
    return wb


def plan_entries():
    return [{'quest': 'Event_Story_1', 'position': position, 'target': 'ded',
             'action_identifier': f'Event_Dialog_{position}'} for position in ('start', 'end')]


class DialoguePipelineTests(unittest.TestCase):
    def test_portraits_face_center_for_all_sides_and_source_directions(self):
        builder = dialogs.load_builder()
        for side, facing, expected in [('left', 'right', None), ('left', 'left', 'horizontal'),
                                       ('right', 'left', None), ('right', 'right', 'horizontal')]:
            with self.subTest(side=side, facing=facing):
                brief = {'id': 100, 'file_name': 'Event_Dialog.proto.js', 'characters': [{
                    'key': 'a', 'name': 'Персонаж', 'side': side,
                    'emotions': {'базовая': 'Portrait'}, 'asset_facing': {'Portrait': facing}}],
                    'lines': [{'speaker': 'a', 'emotion': 'базовая', 'text': 'Привет!'}]}
                _, ui = builder.build(brief)
                frame = ui['params']['dialog'][0]
                portrait = frame['character_' + side]
                self.assertEqual(portrait.get('flip'), expected)
                if expected is None:
                    self.assertNotIn('flip', portrait)
                self.assertNotIn('character_' + ('right' if side == 'left' else 'left'), frame)
                self.assertEqual(frame['speech_bubble']['target_character'], side)

    def test_direction_changes_with_emotion_and_persists_while_silent(self):
        builder = dialogs.load_builder()
        brief = {'id': 100, 'file_name': 'Event_Dialog.proto.js', 'characters': [
            {'key': 'a', 'name': 'Первый', 'side': 'left',
             'emotions': {'базовая': 'A', 'поворот': 'A_Turned'},
             'asset_facing': {'A': 'right', 'A_Turned': 'left'}},
            {'key': 'b', 'name': 'Второй', 'side': 'right',
             'emotions': {'базовая': 'B'}, 'asset_facing': {'B': 'left'}}],
            'lines': [{'speaker': 'a', 'emotion': 'базовая', 'text': 'Раз.'},
                      {'speaker': 'a', 'emotion': 'поворот', 'text': 'Два.'},
                      {'speaker': 'b', 'emotion': 'базовая', 'text': 'Три.'}]}
        _, ui = builder.build(brief)
        frames = ui['params']['dialog']
        self.assertNotIn('flip', frames[0]['character_left'])
        self.assertEqual(frames[1]['character_left']['flip'], 'horizontal')
        self.assertEqual(frames[2]['character_left']['view_class'], 'A_Turned')
        self.assertEqual(frames[2]['character_left']['flip'], 'horizontal')
        self.assertTrue(all('flip' not in f['character_right'] for f in frames))

    def test_unknown_source_direction_is_not_guessed_from_base_emotion(self):
        builder = dialogs.load_builder()
        brief = {'id': 100, 'file_name': 'Event_Dialog.proto.js', 'characters': [{
            'key': 'a', 'name': 'Персонаж', 'side': 'left',
            'emotions': {'базовая': 'A', 'радость': 'A_Happy'},
            'asset_facing': {'A': 'right'}}],
            'lines': [{'speaker': 'a', 'emotion': 'радость', 'text': 'Привет!'}]}
        with self.assertRaisesRegex(ValueError, 'Missing asset_facing for A_Happy'):
            builder.build(brief)
        brief['characters'][0].pop('asset_facing')
        with self.assertRaisesRegex(ValueError, 'needs asset_facing'):
            builder.build(brief)

    def test_pack_passes_asset_directions_to_dialog_builder(self):
        with tempfile.TemporaryDirectory() as tmp:
            plan = {'campaign_id': 'Event', 'pack_id': 'pack_001', 'dialogs': [{
                'quest': 'Event_Story_1', 'position': 'start', 'target': 'ded',
                'characters': [{'key': 'ded', 'side': 'right'}],
                'lines': [{'speaker': 'ded', 'emotion': 'радость', 'text': 'Привет!'}]}]}
            products, _ = dialogs.build_dialogs(plan, registry(), Path(tmp), set(), 100)
            ui = next(obj for obj in products.values() if obj['class'] == 'UserInterfacePrototype')
            portrait = ui['params']['dialog'][0]['character_right']
            self.assertEqual(portrait['view_class'], 'Oseniny_Character_1_Happy')
            self.assertNotIn('flip', portrait)

    def test_insert_and_repeat_preserve_manual_task_data_and_settings(self):
        wb = fixture_quests()
        ws = wb.active
        before = copy.deepcopy(dialogs.scan_quests(ws)['Event_Story_1']['tasks'])
        specs = plan_entries()
        result = dialogs.update_quest_sheet(ws, specs, registry())
        self.assertEqual(result, {'added_tasks': 2, 'existing_tasks': 0})
        tasks = dialogs.scan_quests(ws)['Event_Story_1']['tasks']
        self.assertEqual([t['fields'] for t in tasks[1:-1]], [t['fields'] for t in before])
        self.assertEqual([t['key'] for t in tasks], ['tasks.0', 'tasks.1', 'tasks.2', 'tasks.3'])
        self.assertEqual(ws['D4'].value, 'Oseniny_Character_1')
        self.assertEqual(ws.freeze_panes, 'A5')
        moved = ws.cell(tasks[1]['row'] + 4, 5)
        self.assertTrue(moved.font.bold)
        self.assertEqual(moved.comment.text, 'Ручная правка')
        values = list(ws.values)
        merges = set(map(str, ws.merged_cells.ranges))
        result = dialogs.update_quest_sheet(ws, specs, registry())
        self.assertEqual(result, {'added_tasks': 0, 'existing_tasks': 2})
        self.assertEqual(list(ws.values), values)
        self.assertEqual(set(map(str, ws.merged_cells.ranges)), merges)
        self.assertEqual(dialogs.sheet_snapshot(ws), dialogs.sheet_snapshot(ws))

    def test_duplicate_or_wrong_target_does_not_get_silently_replaced(self):
        wb = fixture_quests()
        specs = plan_entries()
        dialogs.update_quest_sheet(wb.active, specs, registry())
        task = dialogs.scan_quests(wb.active)['Event_Story_1']['tasks'][-1]
        for row in range(task['row'] + 2, task['end']):
            if wb.active.cell(row, 4).value == 'icon':
                wb.active.cell(row, 5, 'AnotherCharacter')
        with self.assertRaisesRegex(ValueError, 'Conflicting existing'):
            dialogs.update_quest_sheet(wb.active, specs, registry())

    def test_portrait_mapping_does_not_change_physical_object_references(self):
        rows = [[None, 'input', 'output', 'helper'], [None, 'donor', 'target', 'Event_Character_1_1'],
                [None, None, None, 'icon', 'Event_Character_1_1']]
        dialogs.map_helper_rows(rows, registry())
        self.assertEqual(rows[1][3], 'Oseniny_Character_1')
        self.assertEqual(rows[2][4], 'Event_Character_1_1')

    def test_bind_preserves_other_behaviours_and_actions_and_is_repeatable(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / 'furniture/Event_Character_1_1.proto.js'
            path.parent.mkdir()
            path.write_bytes(dialogs.json_bytes({'id': 12, 'classname': 'Event_Character_1_1',
                'behaviour': [{'type': 'decoration', 'value': 7},
                              {'type': 'stuff_action', 'actions': ['old_action']}]}))
            wb = Workbook()
            ws = wb.active
            for row in [['ml', 'string', 'string', 'string', 'int'],
                        [None, 'input', 'output', 'classname', 'id'],
                        [None, '/donor.proto.js', '/furniture/Event_Character_1_1.proto.js',
                         'Event_Character_1_1', 12], []]:
                ws.append(row)
            products = dialogs.bind_characters(ws, plan_entries(), registry(), root)
            value = products[path.resolve()]
            self.assertEqual(value['behaviour'][0], {'type': 'decoration', 'value': 7})
            self.assertEqual(value['behaviour'][1]['actions'],
                             ['old_action', 'Event_Dialog_start', 'Event_Dialog_end'])
            values = list(ws.values)
            dialogs.bind_characters(ws, plan_entries(), registry(), root)
            self.assertEqual(list(ws.values), values)

    def test_formula_ids_are_reserved_and_not_written_back(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'manual.xlsx'
            wb = Workbook()
            wb.active.append([1201776, '=A1+1', '=B1+1', 'e100363'])
            wb.save(path)
            before = path.read_bytes()
            ids, tasks = dialogs.occupied_ids([path])
            self.assertTrue({1201776, 1201777, 1201778} <= ids)
            self.assertIn('e100363', tasks)
            self.assertEqual(path.read_bytes(), before)
            self.assertEqual(dialogs.allocate(1201776, ids, '--proto-start'), 1201779)

    def test_task_assignment_skips_ids_reserved_in_other_workbooks(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            wb = fixture_quests()
            dialogs.update_quest_sheet(wb.active, plan_entries(), registry())
            path = root / 'Event_quests.xlsx'
            wb.save(path)
            wb.close()
            result = proto_ids.assign(root, 100, 'e21', backup_dir=root / 'backup',
                                      reserved_proto_ids={100}, reserved_task_ids={'e23'})
            wb = load_workbook(path)
            values = [t['fields']['identifier'] for t in dialogs.scan_quests(wb.active)['Event_Story_1']['tasks']]
            wb.close()
            self.assertEqual(values, ['e24', 'e21', 'e22', 'e25'])
            self.assertEqual(result['next_task_id'], 'e26')
            self.assertEqual(result['next_proto_id'], 101)

    def test_concurrent_edit_aborts_before_any_transaction_write(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            first, second = root / 'first', root / 'second'
            first.write_bytes(b'a')
            second.write_bytes(b'changed')
            with self.assertRaisesRegex(ValueError, 'changed during preparation'):
                dialogs.commit_changes([dialogs.Change(first, b'new', b'a'),
                                        dialogs.Change(second, b'new', b'b')], root / 'backup')
            self.assertEqual(first.read_bytes(), b'a')

    def test_failed_second_replace_rolls_back_first(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            first, second = root / 'first', root / 'second'
            first.write_bytes(b'a')
            second.write_bytes(b'b')
            real_replace = dialogs.os.replace

            def replace(source, target):
                if target == second:
                    raise OSError('locked')
                return real_replace(source, target)

            with patch.object(dialogs.os, 'replace', side_effect=replace):
                with self.assertRaisesRegex(OSError, 'locked'):
                    dialogs.commit_changes([dialogs.Change(first, b'new', b'a'),
                                            dialogs.Change(second, b'new', b'b')], root / 'backup')
            self.assertEqual(first.read_bytes(), b'a')
            self.assertEqual(second.read_bytes(), b'b')

    def test_extra_task_conditions_and_manual_proto_edits_are_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            plan = {'campaign_id': 'Event', 'pack_id': 'pack_001', 'dialogs': [{
                'quest': 'Event_Story_1', 'position': 'start', 'target': 'ded',
                'characters': [{'key': 'ded', 'side': 'left'}],
                'lines': [{'speaker': 'ded', 'emotion': 'базовая', 'text': 'Привет!'}],
                'action': {'conditions': 'active_quest=Event_Story_1+accomplished_task=e1'}}]}
            with self.assertRaisesRegex(ValueError, 'only active_quest'):
                dialogs.build_dialogs(copy.deepcopy(plan), registry(), root, set(), 100)
            plan['dialogs'][0]['action'].pop('conditions')
            products, next_id = dialogs.build_dialogs(plan, registry(), root, set(), 100)
            self.assertEqual(next_id, 102)
            for path, value in products.items():
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_bytes(dialogs.json_bytes(value))
            path = next(p for p in products if '/user_interface/' in p.as_posix())
            value = dialogs.read_json(path)
            value['params']['dialog'][0]['speech_bubble']['texts']['text_tf'] = 'Ручной текст'
            path.write_bytes(dialogs.json_bytes(value))
            with self.assertRaisesRegex(ValueError, 'reconcile manual edits'):
                dialogs.build_dialogs(plan, registry(), root, set(), 100)

    def test_workflow_command_defaults_to_preview_and_routes_parameters(self):
        args = workflow_fast.build_parser().parse_args([
            'dialogs', '--campaign', 'Event', '--pack', 'pack_001', '--game-root', 'game'])
        self.assertFalse(args.apply)
        self.assertIs(args.func, workflow_fast.run_dialogs)
        self.assertIsNone(args.proto_start)


if __name__ == '__main__':
    unittest.main()
