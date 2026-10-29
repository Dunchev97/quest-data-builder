import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from src import export_csv, export_xlsx, technical_quest_tables as tech
from src.validate_proto_exports import expected_updates, table_paths


class TechnicalQuestTablesTests(unittest.TestCase):
    def write_table(self, path, definitions):
        rows = []
        for family, name in definitions:
            output = f"/{family}/Event_2026/{name}.proto.js"
            rows.extend([["sl", "string", "string", "string", "string"],
                         [None, "input", "output", "classname", "ingredients"],
                         [None, "/donor.proto.js", output, name, "asset=Existing_Item:3"], []])
        export_xlsx.write_workbook(path, [export_xlsx.WorkbookSheet("conf", rows, "test")])

    def test_cleanup_tracks_manual_names_all_states_and_deduplicates(self):
        with tempfile.TemporaryDirectory() as tmp:
            a, b = Path(tmp) / "a.xlsx", Path(tmp) / "b.xlsx"
            self.write_table(a, [("quest_item", "Event_2026_Manual_Name"),
                                 ("workbench_item", "Event_2026_FunCollection_4_GR_1"),
                                 ("furniture", "Event_2026_Character_1_2"),
                                 ("chest", "Event_2026_Chest_1_Guest"),
                                 ("debris", "Event_2026_INVITE_Real"),
                                 ("furniture", "Event_2026_Mockery_Sleep"),
                                 ("asset_package", "Event_2026_Package"),
                                 ("quest_action", "Event_2026_Action")])
            self.write_table(b, [("quest_item", "Event_2026_Manual_Name"),
                                 ("recipe", "Event_2026_R_1_Recipe")])
            actions, summary = tech.collect_cleanup("Event_2026", [a, b])
            self.assertEqual((summary["assets"], summary["stuff"], summary["crafting"]), (2, 4, 1))
            self.assertEqual(actions[0], {"do": "cancel_crafting", "workbench": "DayOfDomovoyCaldron",
                                          "recipe_tags": ["Event_2026"]})
            removed = {a["classname"] for a in actions if "classname" in a}
            self.assertEqual(len(removed), 6)
            self.assertIn("Event_2026_Manual_Name", removed)
            self.assertNotIn("Existing_Item", removed)
            self.assertNotIn("Event_2026_Package", removed)
            self.assertTrue(all(set(a) == {"do", "classname"} for a in actions[1:]))

    def test_unknown_families_and_foreign_resources_require_a_classification(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "table.xlsx"
            for family, name, error in [("new_family", "Event_2026_X", "unclassified"),
                                         ("quest_item", "OtherCampaign_GR_1", "foreign")]:
                self.write_table(path, [(family, name)])
                with self.assertRaisesRegex(ValueError, error):
                    tech.collect_cleanup("Event_2026", [path])

    def test_file_contains_group_two_quests_and_two_blank_task_identifiers(self):
        with tempfile.TemporaryDirectory() as tmp:
            campaigns = Path(tmp)
            folder = campaigns / "Event_2026/pack_001/review/proto"
            folder.mkdir(parents=True)
            source = folder / "Event_2026_resources.xlsx"
            self.write_table(source, [("quest_item", "Event_2026_GR_1")])
            before = source.read_bytes()
            result = tech.build("Event_2026", "pack_001", campaigns)
            target = Path(result["xlsx"])
            self.assertEqual(target.name, "Event_2026_technical_quests.xlsx")
            self.assertEqual(source.read_bytes(), before)
            paths = table_paths(target)
            self.assertEqual(len(paths), 5)
            self.assertEqual(len({out for _, _, out in paths}), 3)
            self.assertEqual([out for _, _, out in paths], [
                "/quest_group/event/Event_2026/Event_2026_Tech.proto.js",
                "/quest/event/Event_2026/Event_2026_Tech_1.proto.js",
                "/quest/event/Event_2026/Event_2026_Tech_1.proto.js",
                "/quest/event/Event_2026/Event_2026_Tech_2.proto.js",
                "/quest/event/Event_2026/Event_2026_Tech_2.proto.js"])
            for _, donor, out in paths:
                if donor.startswith("/quest/event/Event_2026/"):
                    self.assertEqual(donor, out)
            updates = expected_updates(target)
            group = updates[0][1]
            self.assertEqual(group["identifier"], "Event_2026_Tech")
            self.assertEqual(group["first_quest"], "Event_2026_Tech_1")
            self.assertEqual(group["last_quest"], "Event_2026_Tech_2")
            self.assertNotIn("spoil_conditions", group)
            tasks = [fields["tasks.0"] for _, fields, _ in updates if "tasks.0" in fields]
            self.assertEqual(len(tasks), 2)
            self.assertTrue(all(t["identifier"] is None for t in tasks))
            clean = updates[3][1]["on_spoil"]
            self.assertEqual(clean, [{"do": "remove_asset", "classname": "Event_2026_GR_1"}])
            tech.validate_cleanup(target, "Event_2026", [source, target])
            wb = load_workbook(target, data_only=False)
            self.assertEqual(wb.sheetnames, ["conf"])
            ws = wb.active
            task_cells = [ws.cell(r, 5) for r in range(1, ws.max_row + 1)
                          if ws.cell(r, 4).value == "identifier" and ws.cell(r - 2, 4).value == "object"]
            self.assertEqual(len(task_cells), 2)
            self.assertTrue(all(c.value is None and c.comment for c in task_cells))
            self.assertIsNone(ws["F3"].comment)  # no task annotation on quest headers
            wb.close()
            snapshot = target.read_bytes()
            with self.assertRaises(FileExistsError):
                tech.build("Event_2026", "pack_001", campaigns)
            self.assertEqual(target.read_bytes(), snapshot)
            self.write_table(source, [("quest_item", "Event_2026_Manual_Rename")])
            with self.assertRaisesRegex(ValueError, "on_spoil missing"):
                tech.validate_cleanup(target, "Event_2026", [source, target])

    def test_cleanup_includes_other_packs_but_never_old_technical_tables(self):
        with tempfile.TemporaryDirectory() as tmp:
            campaigns = Path(tmp)
            for pack in ("pack_001", "pack_002"):
                folder = campaigns / "Event_2026" / pack / "review/proto"
                folder.mkdir(parents=True)
                self.write_table(folder / "Event_2026_resources.xlsx", [("quest_item", f"Event_2026_{pack}")])
            other = campaigns / "Event_2026/pack_002/review/proto" / tech.workbook_name("Event_2026")
            other.write_bytes(b"existing manual technical file is not an inventory source")
            result = tech.build("Event_2026", "pack_001", campaigns,
                                spoil_conditions="time>2026-11-01 23:59")
            self.assertEqual(result["assets"], 2)
            self.assertEqual(expected_updates(Path(result["xlsx"]))[0][1]["spoil_conditions"],
                             "time>2026-11-01 23:59")

    def test_story_export_never_reuses_a_donor_task_identifier(self):
        flattened = dict(export_csv.flatten_task({"identifier": "e9586", "type": "have_location"}))
        self.assertEqual(flattened["identifier"], "")
        self.assertEqual(flattened["type"], "have_location")


if __name__ == "__main__":
    unittest.main()
