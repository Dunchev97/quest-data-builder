import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from src import converter_tables, reviewed_proto_templates as templates
from src.proto_ids import collect_slots
from src.validate_proto_exports import expected_updates, table_paths


class ReviewedProtoTemplatesTests(unittest.TestCase):
    def test_all_audited_references_are_self_referencing_without_production_ids(self):
        config = json.loads(templates.CATALOG.read_text(encoding='utf-8'))
        total = 0
        for template_id, entry in config['reviewed_templates'].items():
            with self.subTest(template_id=template_id):
                path, _ = templates.reference(template_id)
                paths = table_paths(path)
                self.assertEqual(len(paths), entry['prototype_count'])
                self.assertTrue(all(inp == out for row, inp, out in paths))
                wb = load_workbook(path, data_only=False)
                try:
                    protos, tasks = collect_slots([(path, wb)])
                    self.assertTrue(all(cell.value is None for values in protos.values() for _, cell in values))
                    self.assertFalse(tasks)
                    self.assertIsNone(wb.active.freeze_panes)
                finally:
                    wb.close()
                total += len(paths)
        self.assertEqual(total, 168)

    def test_new_prefix_keeps_existing_inputs_and_rewrites_mechanics(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'Fun15_FunCollection_7.xlsx'
            templates.create('FunCollection_7', 'Fun15', path)
            self.assertEqual(converter_tables.validate(path), [])
            self.assertTrue(all('Fun14' in inp and 'Fun15' in out
                                for row, inp, out in table_paths(path)))
            updates = expected_updates(path)
            ingredient = next(fields for output, fields, repl in updates
                              if output.endswith('/Fun15_FunCollection_7_GR_1.proto.js') and output.startswith('/quest_item/'))
            self.assertEqual(ingredient['classname'], 'Fun15_FunCollection_7_GR_1')
            self.assertEqual(ingredient['view_classname'], 'BD23_Resource_GR_Cream')
            self.assertTrue(all(('Fun14', 'Fun15') in repl for output, fields, repl in updates))
            with self.assertRaises(FileExistsError): templates.create('FunCollection_7', 'Fun15', path)

    def test_audited_reward_and_gift_structure_survive_instantiation(self):
        with tempfile.TemporaryDirectory() as tmp:
            good_path = Path(tmp) / 'Fun15_FunCollection_4.xlsx'
            gift_path = Path(tmp) / 'Fun15_FunCollection_5.xlsx'
            templates.create('FunCollection_4', 'Fun15', good_path)
            templates.create('FunCollection_5', 'Fun15', gift_path)
            good = next(fields for output, fields, repl in expected_updates(good_path)
                        if output.endswith('/Fun15_FunCollection_4_REAL.proto.js'))
            rewards = good['extra.stages.1.reward']['one_of']
            self.assertEqual([r['p'] for r in rewards], [20, 20, 20, 20, 15, 15, 3, 3, 2, 2])
            self.assertEqual([r['asset'] for r in rewards], [f'Fun15_FunCollection_4_CL_{i}' for i in range(1, 11)])
            gift = next(fields for output, fields, repl in expected_updates(gift_path)
                        if output.startswith('/quest_item/') and output.endswith('/Fun15_FunCollection_5_GR_2_1.proto.js'))
            self.assertEqual(gift['transformation_to_other_asset_while_gifting.asset'], 'Fun15_FunCollection_5_GR_2_2')

    def test_prefix_replace_precedes_empty_cells_when_reward_array_is_longer(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'Fun15_FunCollection_3.xlsx'
            templates.create('FunCollection_3', 'Fun15', path)
            wb = load_workbook(path, data_only=False)
            try:
                ws = wb.active
                row = next(r for r in range(1, ws.max_row + 1)
                           if ws.cell(r, 3).value == '/chest/Fun/Fun15/Fun15_FunCollection_3.proto.js')
                start = row - 2
                col = next(c for c in range(4, ws.max_column + 1) if ws.cell(start, c).value == 'replace')
                self.assertEqual([ws.cell(row + 1, c).value for c in (col, col + 1)], ['Fun14', 'Fun15'])
                array_col = next(c for c in range(4, ws.max_column + 1) if ws.cell(start, c).value == 'array')
                self.assertIsNotNone(ws.cell(row + 9, array_col).value)
            finally:
                wb.close()


if __name__ == '__main__': unittest.main()
