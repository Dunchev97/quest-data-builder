from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

from src import converter_tables, export_xlsx, proto_ids


class ProtoIdTests(unittest.TestCase):
    def write(self, path, rows):
        export_xlsx.write_workbook(path, [export_xlsx.WorkbookSheet("conf", rows, "fixture")])

    def base(self, output):
        return [["sl", "string", "string", "int", "string"],
                [None, "input", "output", "id", "title"],
                [None, "/donor.proto.js", output, None, "Ручной текст"], []]

    def task(self, output, index):
        return [["ml", "string", "string", "object", None],
                [None, "input", "output", f"tasks.{index}", None],
                [None, output, output, "identifier", None],
                [None, None, None, "type", "action"], []]

    def setup_folder(self, root):
        folder = root / "proto"
        folder.mkdir()
        quest = folder / "Event_quests.xlsx"
        resources = folder / "Event_resources.xlsx"
        q = "/quest/Event_Story_1.proto.js"
        self.write(quest, self.base(q) + self.task(q, 0) + self.task(q, 1))
        # Identical basenames in different prototype families need distinct IDs.
        self.write(resources, self.base("/quest_item/Event_GR_1.proto.js")
                   + self.base("/global_reward/Event_GR_1.proto.js"))
        return folder, quest, resources

    def test_assigns_once_per_output_and_preserves_manual_content_and_metadata(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            folder, quest, resources = self.setup_folder(root)
            wb = load_workbook(quest)
            ws = wb.active
            ws["E3"] = "Исправлено вручную: домовёнок"
            ws["E3"].font = Font(bold=True, color="AA1122")
            ws["E3"].comment = Comment("Ручной комментарий", "Автор")
            ws.column_dimensions["E"].width = 48
            ws.row_dimensions[3].height = 45
            ws.freeze_panes = "A16"
            merges = set(map(str, ws.merged_cells.ranges))
            wb.save(quest)
            wb.close()
            result = proto_ids.assign(folder, 500, "e800", backup_dir=root / "backup")
            self.assertEqual((result["prototypes"], result["tasks"]), (3, 2))
            self.assertEqual((result["proto_first"], result["proto_last"]), (500, 502))
            self.assertEqual((result["next_proto_id"], result["next_task_id"]), (503, "e802"))
            ids, tasks = proto_ids.read_assigned_ids([quest, resources])
            self.assertEqual(ids["/quest/Event_Story_1.proto.js"], 500)
            self.assertEqual(set(tasks.values()), {"e800", "e801"})
            self.assertEqual(converter_tables.validate(quest, assigned_ids=True), [])
            self.assertTrue(converter_tables.validate(quest))  # authoring mode stays strict
            wb = load_workbook(quest)
            ws = wb.active
            self.assertEqual(ws["E3"].value, "Исправлено вручную: домовёнок")
            self.assertEqual(ws["E3"].comment.text, "Ручной комментарий")
            self.assertEqual(ws.column_dimensions["E"].width, 48)
            self.assertEqual(ws.row_dimensions[3].height, 45)
            self.assertEqual(ws.freeze_panes, "A16")
            self.assertEqual(set(map(str, ws.merged_cells.ranges)), merges)
            wb.close()
            self.assertEqual(len(list(folder.iterdir())), 2)

    def test_repeat_is_idempotent_and_never_overwrites_assigned_values(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            folder, quest, resources = self.setup_folder(root)
            proto_ids.assign(folder, 500, "e800", backup_dir=root / "backup")
            before = {p: p.read_bytes() for p in folder.glob("*.xlsx")}
            result = proto_ids.assign(folder, 500, "e800", backup_dir=root / "unused")
            self.assertEqual(result["changed_files"], 0)
            self.assertEqual((result["next_proto_id"], result["next_task_id"]), (503, "e802"))
            self.assertTrue(all(p.read_bytes() == content for p, content in before.items()))
            self.assertFalse((root / "unused").exists())

    def test_existing_values_are_reserved_before_filling_blanks(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            folder, quest, resources = self.setup_folder(root)
            wb = load_workbook(resources)
            wb.active["D3"] = 500
            wb.save(resources)
            wb.close()
            proto_ids.assign(folder, 500, "e800", backup_dir=root / "backup")
            ids, _ = proto_ids.read_assigned_ids([quest, resources])
            self.assertEqual(ids["/quest_item/Event_GR_1.proto.js"], 500)
            self.assertEqual(ids["/quest/Event_Story_1.proto.js"], 501)
            self.assertEqual(len(set(ids.values())), 3)

    def test_duplicate_ids_are_rejected_before_any_canonical_file_changes(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            folder, quest, resources = self.setup_folder(root)
            for path in (quest, resources):
                wb = load_workbook(path)
                wb.active["D3"] = 600
                wb.save(path)
                wb.close()
            before = {p: p.read_bytes() for p in (quest, resources)}
            with self.assertRaisesRegex(ValueError, "Duplicate proto id"):
                proto_ids.assign(folder, 500, "e800", backup_dir=root / "backup")
            self.assertTrue(all(p.read_bytes() == v for p, v in before.items()))

    def test_excel_lock_stops_the_batch_without_partial_assignments(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            folder, quest, resources = self.setup_folder(root)
            before = {p: p.read_bytes() for p in (quest, resources)}
            with patch.object(proto_ids, "check_replace_access", side_effect=PermissionError("Excel lock")):
                with self.assertRaises(PermissionError):
                    proto_ids.assign(folder, 500, "e800", backup_dir=root / "backup")
            self.assertTrue(all(p.read_bytes() == v for p, v in before.items()))

    def test_missing_or_invalid_assigned_ids_and_duplicate_task_ids_fail(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            folder, quest, resources = self.setup_folder(root)
            self.assertTrue(proto_ids.assigned_id_errors([quest, resources]))
            proto_ids.assign(folder, 500, "e800", backup_dir=root / "backup")
            wb = load_workbook(quest)
            task_cells = [c for cells in proto_ids.scan_sheet(wb.active)[1].values() for c in cells]
            task_cells[1].value = task_cells[0].value
            wb.save(quest)
            wb.close()
            self.assertIn("Duplicate task identifier", proto_ids.assigned_id_errors([quest, resources])[0])
            wb = load_workbook(resources)
            wb.active["D3"] = "500"
            wb.save(resources)
            wb.close()
            self.assertIn("invalid proto id", proto_ids.assigned_id_errors([resources])[0])

    def test_duplicate_id_columns_of_one_prototype_share_a_single_allocation(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            folder = root / "proto"
            folder.mkdir()
            path = folder / "recipe.xlsx"
            rows = [["sl", "string", "string", "int", "int"],
                    [None, "input", "output", "id", "id"],
                    [None, "/old.proto.js", "/new.proto.js", None, None]]
            self.write(path, rows)
            result = proto_ids.assign(folder, 500, "e800", backup_dir=root / "backup")
            self.assertEqual(result["prototypes"], 1)
            self.assertEqual(result["next_proto_id"], 501)
            self.assertEqual(converter_tables.validate(path, assigned_ids=True), [])
            wb = load_workbook(path)
            self.assertEqual([wb.active.cell(3, c).value for c in (4, 5)], [500, 500])
            wb.close()


if __name__ == "__main__":
    unittest.main()
