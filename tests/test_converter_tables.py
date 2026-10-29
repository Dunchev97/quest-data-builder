import json
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from src import converter_tables, export_xlsx, fun_collection_tables
from src.validate_proto_exports import expected_updates, replacement_audit_text, run, table_paths
from src.workflow_context import file_sha256


class ConverterTablesTests(unittest.TestCase):
    def test_replacement_audit_keeps_only_the_explicit_graphic_alias_exception(self):
        data = {'classname': 'Fun15_Item', 'view_classname': 'Fun13_Art',
                'nested': [{'view_classname': 'Fun12_Art', 'asset': 'Fun13_WrongResource'}]}
        audited = replacement_audit_text(data)
        self.assertNotIn('Fun13_Art', audited)
        self.assertNotIn('Fun12_Art', audited)
        self.assertIn('Fun13_WrongResource', audited)
        self.assertEqual(data['view_classname'], 'Fun13_Art')

    def test_replace_after_a_gap_is_rejected_even_inside_a_long_reward_array(self):
        wb = Workbook()
        ws = wb.active
        for row in [['ml', 'string', 'string', 'array', 'replace', None],
                    [None, 'input', 'output', 'rand_reward.one_of', 'find', 'replace'],
                    [None, '/a.proto.js', '/b.proto.js', '{"asset":"One"}', 'Old', 'New'],
                    [None, None, None, '{"asset":"Two"}', None, None],
                    [None, None, None, None, 'Source', 'Target']]:
            ws.append(row)
        converter_tables.prepare_sheet(ws)
        errors = converter_tables.validate_sheet(ws)
        self.assertTrue(any('replace entry after a gap' in e for e in errors))

    def test_standalone_fun_outputs_cannot_keep_a_foreign_donor_directory(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'Fun14_FunCollection_7.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.title = 'conf'
            ws.append(['sl', 'string', 'string', 'string'])
            ws.append([None, 'input', 'output', 'classname'])
            donor = '/asset_package/HeavenlyExotic/Fun13_Item_Package.proto.js'
            ws.append([None, donor, '/asset_package/HeavenlyExotic/Fun14_Item_Package.proto.js',
                       'Fun14_Item_Package'])
            wb.save(path)
            self.assertTrue(any('output directory' in e for e in converter_tables.validate(path)))
            for output in ('/asset_package/Fun14/Fun14_Item_Package.proto.js',
                           '/asset_package/Fun/Fun14/Fun14_Item_Package.proto.js',
                           '/post_action/ask_for_Fun14_Item.proto.js'):
                with self.subTest(output=output):
                    ws['C3'] = output
                    wb.save(path)
                    self.assertEqual(converter_tables.validate(path), [])
                    self.assertEqual(ws['B3'].value, donor)
            ws['C3'] = '/post_action/ask_for_Fun13_Item.proto.js'
            wb.save(path)
            self.assertTrue(any('campaign prefix' in e for e in converter_tables.validate(path)))
            wb.close()

    def test_fun7_package_template_outputs_use_the_target_event_directory(self):
        templates = json.loads(fun_collection_tables.TEMPLATES.read_text(encoding='utf-8'))['templates']
        template = next(t for t in templates if t['id'] == 'FunCollection_7')
        packages = [row for block in template['blocks'] for row in block.get('rows', [])
                    if row.get('output', '').startswith('/asset_package/')]
        self.assertEqual(len(packages), 3)
        for row in packages:
            output = row['output'].replace('{prefix}', 'Fun14')
            self.assertEqual(Path(output).parent.as_posix(), '/asset_package/Fun14')
            self.assertTrue(row['input'].startswith('/asset_package/Fun/Fun10/'))

    def test_task_cannot_read_a_stale_path_after_quest_output_is_moved(self):
        wb = Workbook()
        ws = wb.active
        for row in [["ml", "string", "string", "object", None],
                    [None, "input", "output", "tasks.0", None],
                    [None, "/quest/Old/tech/Old_Tech_1.proto.js",
                     "/quest/event/Old/Old_Tech_1.proto.js", "identifier", None]]:
            ws.append(row)
        converter_tables.prepare_sheet(ws)
        errors = converter_tables.validate_sheet(ws)
        self.assertTrue(any("task input/output must both reference" in e for e in errors))

    def test_each_entity_and_each_replaced_row_gets_its_own_block(self):
        rows = [
            ['ml','string','string','string'], ['', 'input','output','classname'],
            ['', '/furniture/a.proto.js','/furniture/x.proto.js','x'],
            ['', '/furniture/a.proto.js','/furniture/y.proto.js','y'], [],
            ['sl','string','string','string','replace',''],
            ['', 'input','output','identifier','find','replace'],
            ['', '/action/a.proto.js','/action/x.proto.js','x','Old','New'],
            ['', '/action/a.proto.js','/action/y.proto.js','y','Old','New'],
        ]
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'actions.xlsx'
            export_xlsx.write_workbook(path,[export_xlsx.WorkbookSheet('conf',rows,'generated_actions')])
            self.assertEqual(converter_tables.validate(path),[])
            wb=load_workbook(path)
            self.assertEqual([r[0].value for r in wb.active if r[0].value],['ml','ml','sl','sl'])
            self.assertEqual(len(table_paths(path)),4)
            # There must be no stored empty-string cells: JS uses null as the sentinel.
            with ZipFile(path) as archive:
                xml=archive.read('xl/worksheets/sheet1.xml').decode()
            self.assertNotIn('t="inlineStr"></c>',xml)

    def test_task_and_replace_pairs_are_merged_and_values_preserved(self):
        rows=[['ml','string','string','object'],['','input','output','tasks.0',''],
              ['', '/quest/x.proto.js','/quest/x.proto.js','type','get_asset'],
              ['', '', '', 'amount', 5], ['', '', '', 'hint', 'Спасибо, дружок!']]
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'quests.xlsx'
            export_xlsx.write_workbook(path,[export_xlsx.WorkbookSheet('conf',rows,'generated_quests')])
            self.assertEqual(converter_tables.validate(path),[])
            wb=load_workbook(path)
            self.assertEqual({str(m) for m in wb.active.merged_cells.ranges},{'D1:E1','D2:E2'})
            intended=expected_updates(path)[0][1]['tasks.0']
            self.assertEqual(intended,{'type':'get_asset','amount':5,'hint':'Спасибо, дружок!'})

    def test_invalid_numeric_and_duplicate_fields_fail_before_saving(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'bad.xlsx'
            rows=[['sl','string','string','int'],['','input','output','price'],['','/a.proto.js','/b.proto.js','oops']]
            with self.assertRaisesRegex(ValueError,'expected a number'):
                export_xlsx.write_workbook(path,[export_xlsx.WorkbookSheet('conf',rows,'interactive')])
            self.assertFalse(path.exists())
            rows=[['sl','string','string','string','string'],['','input','output','title','title'],['','/a.proto.js','/b.proto.js','one','two']]
            with self.assertRaisesRegex(ValueError,'duplicate field'):
                export_xlsx.write_workbook(path,[export_xlsx.WorkbookSheet('conf',rows,'interactive')])

    def test_validator_checks_ids_in_all_data_rows_and_rejects_old_directives(self):
        wb=Workbook(); ws=wb.active; ws.title='conf'
        for row in [['sl','string','string','int'],['','input','output','id'],['','/a.proto.js','/b.proto.js',None],['','/a.proto.js','/c.proto.js',7]]:
            ws.append(row)
        self.assertTrue(any('id must be blank' in e for e in converter_tables.validate_sheet(ws)))
        ws['A1']='temp_01'
        self.assertTrue(any('unsupported directive' in e for e in converter_tables.validate_sheet(ws)))

    def test_legacy_button_price_and_array_dot_paths_are_repaired(self):
        rows=[['temp_01','string','string','string','int','int'],
              ['','input','output','classname','extra.open_now_btn_window','price'],
              ['','/a.proto.js','/b.proto.js','B','Получить за {price}','15']]
        out=converter_tables.normalize_rows(rows)
        self.assertEqual(out[0],['sl','string','string','string','string','int'])
        self.assertEqual(out[2][-1],15)
        rows=[['sl','string','string','string'],['','input','output','behaviour.description_1'],['','/a.proto.js','/b.proto.js','Текст']]
        self.assertEqual(converter_tables.normalize_rows(rows)[1][3],'behaviour.0.window_spec.description_1')

    def test_split_preserves_master_and_manual_text_and_protects_exports(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'master.xlsx'
            wb=Workbook(); ws=wb.active; ws.title='КВЕСТЫ'; ws.freeze_panes='A4'
            for row in [['sl','string','string','string'],['','input','output','title'],['','/a.proto.js','/b.proto.js','Ручная правка']]: ws.append(row)
            wb.save(path)
            original=path.read_bytes()
            result=export_xlsx.split_stage6_workbook(path,campaign_id='Event_2026')
            target=Path(result['files'][0]['xlsx'])
            conf=load_workbook(target)
            self.assertEqual(conf.sheetnames,['conf'])
            self.assertEqual(target.name,'Event_2026_quests.xlsx')
            self.assertEqual(conf.active['D3'].value,'Ручная правка')
            self.assertIsNone(conf.active.freeze_panes)
            self.assertEqual(path.read_bytes(),original)
            self.assertFalse((target.parent/'manifest.json').exists())
            with self.assertRaises(FileExistsError): export_xlsx.split_stage6_workbook(path,campaign_id='Event_2026')

    def test_fun4_exact_family_has_22_prototypes_and_distinct_window_fields(self):
        selection=json.loads((fun_collection_tables.ROOT/'campaigns/VelesNight_2026/fun_interactive_objects.selection.json').read_text(encoding='utf-8'))['selected_objects'][0]
        rows=fun_collection_tables.build_rows(selection,'VelesNight_2026')
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'funcollection_4.xlsx'
            export_xlsx.write_workbook(path,[export_xlsx.WorkbookSheet('conf',rows,'fun:FunCollection_4')])
            self.assertEqual(converter_tables.validate(path),[])
            paths=table_paths(path)
            self.assertEqual(len(paths),22)
            outputs=[p[2] for p in paths]
            self.assertEqual(sum('_CL_' in p for p in outputs),10)
            self.assertTrue(any('/VelesNight_2026_INVITE_Real.proto.js' in p for p in outputs))
            self.assertFalse(any('_INVITE_CL_' in p or '_INVITE_GR_' in p for p in outputs))
            writes=expected_updates(path)
            good=next(fields for output,fields,repl in writes if output.endswith('_INVITE_Real.proto.js'))
            self.assertIn('extra.pop_up_message',good)
            self.assertIn('extra.description_window',good)
            self.assertNotEqual(good['extra.pop_up_message'],good['extra.description_window'])

    def test_dry_run_rejects_the_real_game_tree_as_sandbox(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            with self.assertRaisesRegex(ValueError,'separate'):
                run([],root/'converter.js',root,root/'output')

    def test_review_folder_checksum_tracks_all_xlsx_and_ignores_excel_locks(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            (root/'quests.xlsx').write_bytes(b'quests')
            (root/'actions.xlsx').write_bytes(b'actions')
            before=file_sha256(root)
            (root/'~$quests.xlsx').write_bytes(b'lock')
            self.assertEqual(file_sha256(root),before)
            (root/'actions.xlsx').write_bytes(b'manual edit')
            self.assertNotEqual(file_sha256(root),before)

    def test_fun4_default_output_is_in_pack_proto_and_existing_edits_survive(self):
        with tempfile.TemporaryDirectory() as tmp:
            campaigns=Path(tmp)/'campaigns'
            campaign=campaigns/'VelesNight_2026'
            campaign.mkdir(parents=True)
            source=fun_collection_tables.ROOT/'campaigns/VelesNight_2026/fun_interactive_objects.selection.json'
            (campaign/source.name).write_bytes(source.read_bytes())
            result=fun_collection_tables.build_campaign('VelesNight_2026','pack_001',campaigns)
            target=campaign/'pack_001/review/proto/VelesNight_2026_funcollection_4.xlsx'
            self.assertEqual(Path(result['xlsx']),target)
            self.assertFalse((campaign/'funcollection_4.xlsx').exists())
            before=target.read_bytes()
            fun_collection_tables.ensure_selected_tables('VelesNight_2026','pack_001',campaigns)
            self.assertEqual(target.read_bytes(),before)
            self.assertEqual({p.name for p in target.parent.iterdir()},
                             {target.name, 'VelesNight_2026_shop_invisible_packages.txt',
                              'VelesNight_2026_DayOfDomovoyCaldron.txt'})

    def test_campaign_filenames_include_prefix_without_duplicating_it(self):
        self.assertEqual(export_xlsx.campaign_workbook_name('Event_2026','quests.xlsx'),'Event_2026_quests.xlsx')
        existing='quest_Event_2026_accomplished.xlsx'
        self.assertEqual(export_xlsx.campaign_workbook_name('Event_2026',existing),existing)
        with self.assertRaises(ValueError):
            export_xlsx.campaign_workbook_name('','quests.xlsx')


if __name__=='__main__':
    unittest.main()
